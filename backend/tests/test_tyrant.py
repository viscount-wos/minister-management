"""Frost Dragon Tyrant: answers, settings (windows), profile split, admin list/summary/exports."""
import csv
import io

import openpyxl

from tests.conftest import start_round

EVENT = 'tyrant'
TROOPS = {'infantry': {'furnace_level': 'FC5', 'tier': 10}, 'lancer': {'furnace_level': 'FC3', 'tier': 9},
          'marksman': {'furnace_level': None, 'tier': 11}}


def put(client, fid, profile=None, answers=None):
    prof = {'game_name': f'P{fid}', 'alliance': 'abc'}
    if profile:
        prof.update(profile)
    return client.put(f'/api/events/{EVENT}/current/application/{fid}', json={'profile': prof, 'answers': answers or {}})


def good_answers(**kw):
    a = {'availability': ['w1', 'w3'], 'discord_vc': True, 'gem_spend': 10000, 'roles': ['joiner', 'rally_leader'],
         'language': 'en'}
    a.update(kw)
    return a


def test_default_windows_and_public_round(client, admin):
    start_round(client, admin, 'FDT 1', event=EVENT)
    r = client.get('/api/events/tyrant/current')
    assert r.status_code == 200
    w = r.json['settings']['windows']
    assert [(x['id'], x['start'], x['end'], x['rush']) for x in w] == [
        ('w1', '11:01', '11:15', True), ('w2', '11:15', '13:00', False), ('w3', '13:00', '15:00', False),
        ('w4', '15:00', '16:30', False), ('w5', '16:30', '18:00', False)]


def test_submit_profile_and_round_split(client, admin):
    start_round(client, admin, 'FDT 1', event=EVENT)
    r = put(client, '1001', profile={'discord_id': 'alice#1234', 'furnace_level': 'fc5', 'power': 410_500_000,
                                     'troops': TROOPS},
            answers=good_answers())
    assert r.status_code == 201, r.json
    prof = r.json['profile']
    assert prof['alliance'] == 'ABC' and prof['discord_id'] == 'alice#1234' and prof['furnace_level'] == 'FC5'
    assert prof['power'] == 410_500_000 and prof['troops']['marksman'] == {'furnace_level': None, 'tier': 11}
    ans = r.json['application']['answers']
    # canonical order (windows order / role list order), not the order sent
    assert ans == {'availability': ['w1', 'w3'], 'discord_vc': True, 'gem_spend': 10000,
                   'roles': ['rally_leader', 'joiner'], 'language': 'en'}
    # profile is shared: visible through the profile endpoint, reusable by other events
    p = client.get('/api/profile/1001').json
    assert p['discord_id'] == 'alice#1234' and p['troops']['infantry'] == {'furnace_level': 'FC5', 'tier': 10}


def test_requires_name_and_alliance(client, admin):
    start_round(client, admin, 'FDT', event=EVENT)
    r = client.put('/api/events/tyrant/current/application/5', json={'profile': {'game_name': 'X'}, 'answers': {}})
    assert r.status_code == 400 and r.json['field'] == 'profile.alliance'
    r = put(client, '5', profile={'alliance': 'ABCD'})
    assert r.status_code == 400 and r.json['field'] == 'profile.alliance'


def test_blank_answers_are_valid(client, admin):
    start_round(client, admin, 'FDT', event=EVENT)
    r = put(client, '6')
    assert r.status_code == 201
    assert r.json['application']['answers'] == {'availability': [], 'discord_vc': False, 'gem_spend': None,
                                                 'roles': [], 'language': None}


def test_strict_answer_validation(client, admin):
    start_round(client, admin, 'FDT', event=EVENT)
    cases = [
        ({'availability': ['w9']}, 'answers.availability'),
        ({'availability': 'w1'}, 'answers.availability'),
        ({'roles': ['king']}, 'answers.roles'),
        ({'discord_vc': 'maybe'}, 'answers.discord_vc'),
        ({'gem_spend': -1}, 'answers.gem_spend'),
        ({'gem_spend': 1.5}, 'answers.gem_spend'),
        ({'gem_spend': 'lots'}, 'answers.gem_spend'),
        ({'gem_spend': 10 ** 10}, 'answers.gem_spend'),
        ({'language': 'xx'}, 'answers.language'),
        ({'surprise': 1}, 'answers.surprise'),
    ]
    for answers, field in cases:
        r = put(client, '7', answers=answers)
        assert r.status_code == 400 and r.json['code'] == 'VALIDATION_ERROR' and r.json['field'] == field, (answers, r.json)
    assert client.get('/api/events/tyrant/current/application/7').status_code == 404  # nothing written


def test_strict_troop_validation(client, admin):
    start_round(client, admin, 'FDT', event=EVENT)
    cases = [
        ({'troops': {'archer': {}}}, 'profile.troops'),
        ({'troops': {'infantry': {'furnace_level': 'FC11'}}}, 'profile.troops.infantry.furnace_level'),
        ({'troops': {'infantry': {'furnace_level': 31}}}, 'profile.troops.infantry.furnace_level'),
        ({'troops': {'infantry': {'tier': 12}}}, 'profile.troops.infantry.tier'),
        ({'troops': {'infantry': {'tier': 'T10'}}}, 'profile.troops.infantry.tier'),
        ({'troops': {'infantry': {'level': 3}}}, 'profile.troops.infantry'),
        ({'troops': [1, 2]}, 'profile.troops'),
        ({'furnace_level': 'FC11'}, 'profile.furnace_level'),
        ({'furnace_level': '31'}, 'profile.furnace_level'),
        ({'furnace_level': 0}, 'profile.furnace_level'),
        # Tyrant: Fire Crystal levels only (owner rule p2d), for the furnace and every camp
        ({'furnace_level': '30'}, 'profile.furnace_level'),
        ({'furnace_level': 25}, 'profile.furnace_level'),
        ({'troops': {'lancer': {'furnace_level': '28', 'tier': 9}}}, 'profile.troops.lancer.furnace_level'),
        ({'troops': {'marksman': {'furnace_level': 1}}}, 'profile.troops.marksman.furnace_level'),
        ({'power': -5}, 'profile.power'),
        ({'discord_id': 'x' * 65}, 'profile.discord_id'),
    ]
    for prof, field in cases:
        r = put(client, '8', profile=prof)
        assert r.status_code == 400 and r.json['field'] == field, (prof, r.json)
    r = put(client, '8', profile={'troops': {'infantry': {'furnace_level': 'FC2'}}})
    assert r.status_code == 201
    blank = {'furnace_level': None, 'tier': None}
    assert r.json['profile']['troops'] == {'infantry': {'furnace_level': 'FC2', 'tier': None}, 'lancer': blank,
                                           'marksman': blank}


def test_windows_settings_validation_and_carry_over(client, admin):
    rnd = start_round(client, admin, 'FDT', event=EVENT)
    url = f'/api/admin/rounds/{rnd["id"]}'
    bad = [
        ({'windows': []}, 'settings.windows'),
        ({'windows': [{'id': 'a', 'start': '12:00', 'end': '11:00'}]}, 'settings.windows.0.end'),
        ({'windows': [{'id': 'a', 'start': '1200', 'end': '13:00'}]}, 'settings.windows.0.start'),
        ({'windows': [{'id': 'A B', 'start': '10:00', 'end': '11:00'}]}, 'settings.windows.0.id'),
        ({'windows': [{'id': 'a', 'start': '10:00', 'end': '11:00'}, {'id': 'a', 'start': '12:00', 'end': '13:00'}]},
         'settings.windows.1.id'),
        ({'windows': [{'id': 'a', 'start': '10:00', 'end': '11:00', 'colour': 'red'}]}, 'settings.windows.0'),
        ({'research_day': 'friday'}, 'settings'),
    ]
    for settings, field in bad:
        r = client.put(url, json={'settings': settings}, headers=admin)
        assert r.status_code == 400 and r.json['field'] == field, (settings, r.json)
    new = [{'id': 'late', 'start': '19:00', 'end': '20:00'}, {'id': 'rush', 'start': '10:00', 'end': '10:15', 'rush': True}]
    r = client.put(url, json={'settings': {'windows': new}, 'closing_time': '2099-01-01T00:00:00Z'}, headers=admin)
    assert r.status_code == 200
    assert [w['id'] for w in r.json['settings']['windows']] == ['rush', 'late']  # sorted by start
    assert r.json['closing_time'] == '2099-01-01T00:00:00Z'
    assert client.get('/api/events/tyrant/current').json['settings']['windows'][0] == \
        {'id': 'rush', 'start': '10:00', 'end': '10:15', 'rush': True}
    r2 = start_round(client, admin, 'FDT 2', event=EVENT)
    assert [w['id'] for w in r2['settings']['windows']] == ['rush', 'late']  # carried over


def test_removed_window_kept_on_resubmit(client, admin):
    rnd = start_round(client, admin, 'FDT', event=EVENT)
    assert put(client, '9', answers=good_answers(availability=['w2', 'w5'])).status_code == 201
    client.put(f'/api/admin/rounds/{rnd["id"]}', headers=admin,
               json={'settings': {'windows': [{'id': 'w2', 'start': '11:15', 'end': '13:00'}]}})
    r = put(client, '9', answers=good_answers(availability=['w2', 'w5']))
    assert r.status_code == 200 and r.json['application']['answers']['availability'] == ['w2', 'w5']
    r = put(client, '10', answers=good_answers(availability=['w5']))
    assert r.status_code == 400 and r.json['field'] == 'answers.availability'


def test_new_vs_edit_and_previous_round(client, admin):
    start_round(client, admin, 'FDT 1', event=EVENT)
    assert put(client, '11', answers=good_answers()).status_code == 201
    assert put(client, '11', answers=good_answers(gem_spend=5)).status_code == 200
    r2 = start_round(client, admin, 'FDT 2', event=EVENT)
    assert client.get('/api/events/tyrant/current/application/11').status_code == 404  # new round: new application
    prev = client.get('/api/events/tyrant/previous-application/11')
    assert prev.status_code == 200 and prev.json['round_name'] == 'FDT 1' and prev.json['answers']['gem_spend'] == 5
    assert client.get('/api/events/ministry/previous-application/11').status_code == 404
    assert client.get('/api/profile/11').json['alliance'] == 'ABC'  # profile pre-fills
    assert r2['status'] == 'open'


def test_closing_time_blocks_new_only(client, admin):
    rnd = start_round(client, admin, 'FDT', event=EVENT)
    assert put(client, '12').status_code == 201
    client.put(f'/api/admin/rounds/{rnd["id"]}', json={'closing_time': '2000-01-01T00:00:00Z'}, headers=admin)
    r = put(client, '13')
    assert r.status_code == 403 and r.json['code'] == 'APPLICATIONS_CLOSED'
    assert put(client, '12', answers={'discord_vc': True}).status_code == 200


def _seed(client, admin):
    rnd = start_round(client, admin, 'FDT', event=EVENT)
    put(client, '21', profile={'game_name': 'Zed', 'alliance': 'AAA', 'furnace_level': 'FC10', 'power': 900_000_000,
                               'troops': TROOPS, 'discord_id': 'zed'},
        answers=good_answers(availability=['w1', 'w2'], roles=['rally_leader'], gem_spend=50000))
    # 22 has a LEGACY pre-FC furnace '30' (stored through the generic profile route, as Minister / pre-p2d tyrant
    # sign-ups did); the tyrant submit doesn't resend it, so it stays as stored
    assert client.put('/api/admin/profiles/22', json={'game_name': 'amy', 'furnace_level': '30'},
                      headers=admin).status_code in (200, 201)
    put(client, '22', profile={'game_name': 'amy', 'alliance': 'BBB', 'power': 100_000_000,
                               'troops': {'infantry': {'furnace_level': 'FC1', 'tier': 8}}},
        answers=good_answers(availability=['w3'], discord_vc=False, roles=['joiner', 'gathering'], gem_spend=None))
    put(client, '23', profile={'game_name': '=HYPERLINK("x")', 'alliance': 'aaa'},
        answers=good_answers(availability=['w1'], roles=[], gem_spend=0))
    return rnd


def test_admin_list_filter_search_sort_paging(client, admin):
    rnd = _seed(client, admin)
    base = f'/api/admin/tyrant/rounds/{rnd["id"]}/applications'
    assert client.get(base).status_code == 401
    r = client.get(base, headers=admin)
    assert r.status_code == 200 and r.json['total'] == 3
    assert [a['fid'] for a in r.json['applications']] == ['23', '22', '21']  # newest first
    assert r.json['applications'][0]['profile']['game_name'] == '=HYPERLINK("x")'
    r = client.get(base + '?alliance=aaa', headers=admin)
    assert sorted(a['fid'] for a in r.json['applications']) == ['21', '23']
    r = client.get(base + '?q=AM', headers=admin)
    assert [a['fid'] for a in r.json['applications']] == ['22']
    r = client.get(base + '?q=zed', headers=admin)  # discord id / name
    assert [a['fid'] for a in r.json['applications']] == ['21']
    r = client.get(base + '?sort=power&dir=desc', headers=admin)
    assert [a['fid'] for a in r.json['applications']] == ['21', '22', '23']  # blank power last
    r = client.get(base + '?sort=furnace&dir=desc', headers=admin)
    assert [a['fid'] for a in r.json['applications']] == ['21', '22', '23']  # FC10 > 30 > blank
    r = client.get(base + '?sort=furnace&dir=asc', headers=admin)
    assert [a['fid'] for a in r.json['applications']] == ['22', '21', '23']
    r = client.get(base + '?sort=name&dir=asc', headers=admin)
    assert [a['fid'] for a in r.json['applications']] == ['23', '22', '21']
    r = client.get(base + '?sort=gems&dir=asc&limit=1&offset=1', headers=admin)
    assert r.json['total'] == 3 and r.json['limit'] == 1 and [a['fid'] for a in r.json['applications']] == ['21']
    r = client.get(base + '?min_furnace=fc1', headers=admin)
    assert [a['fid'] for a in r.json['applications']] == ['21']  # FC10 yes, 30 and blank no
    r = client.get(base + '?min_furnace=20', headers=admin)
    assert sorted(a['fid'] for a in r.json['applications']) == ['21', '22']
    assert client.get(base + '?min_furnace=FC11', headers=admin).json['field'] == 'min_furnace'
    assert client.get(base + '?sort=bogus', headers=admin).json['field'] == 'sort'
    assert client.get(base + '?limit=0', headers=admin).json['field'] == 'limit'
    r = client.get('/api/admin/tyrant/rounds/current/applications', headers=admin)
    assert r.json['round_id'] == rnd['id']
    m = start_round(client, admin, 'Ministry', event='ministry')
    assert client.get(f'/api/admin/tyrant/rounds/{m["id"]}/applications', headers=admin).status_code == 404
    # generic list also works for tyrant (MCP list_applications)
    g = client.get(f'/api/admin/rounds/{rnd["id"]}/applications', headers=admin)
    assert g.status_code == 200 and g.json['total'] == 3


def test_admin_summary(client, admin):
    rnd = _seed(client, admin)
    s = client.get(f'/api/admin/tyrant/rounds/{rnd["id"]}/summary', headers=admin).json
    assert s['total'] == 3 and s['opening_rush'] == 2 and s['discord_vc'] == 2
    assert 'gem_spend_total' not in s  # owner: no aggregate gem total anywhere
    assert s['round_total'] == 3 and s['alliance_options'] == ['AAA', 'BBB']
    assert {w['id']: w['count'] for w in s['windows']} == {'w1': 2, 'w2': 1, 'w3': 1, 'w4': 0, 'w5': 0}
    assert s['alliances'] == [{'alliance': 'AAA', 'count': 2}, {'alliance': 'BBB', 'count': 1}]
    assert s['roles'] == {'rally_leader': 1, 'joiner': 1, 'gathering': 1, 'battle_mgmt': 0, 'event_prep': 0}
    assert s['troop_tiers']['infantry'] == {'T10': 1, 'T8': 1, 'none': 1}
    assert s['troop_tiers']['marksman'] == {'T11': 1, 'none': 2}
    assert s['furnace_levels'] == {'FC10': 1, '30': 1, 'none': 1}  # FC above pre-FC, blanks last
    s = client.get(f'/api/admin/tyrant/rounds/{rnd["id"]}/summary?alliance=BBB', headers=admin).json
    assert s['total'] == 1 and s['opening_rush'] == 0 and s['round_total'] == 3
    assert s['alliance_options'] == ['AAA', 'BBB']  # the picker keeps every alliance of the round


def test_exports_are_formula_safe(client, admin):
    rnd = _seed(client, admin)
    r = client.get(f'/api/admin/tyrant/rounds/{rnd["id"]}/export.csv', headers=admin)
    assert r.status_code == 200 and r.mimetype == 'text/csv' and 'attachment' in r.headers['Content-Disposition']
    rows = list(csv.reader(io.StringIO(r.data.decode('utf-8-sig'))))
    head = rows[0]
    assert head[:4] == ['FID', 'In-Game Name', 'Alliance', 'Discord ID'] and '11:01-11:15 Opening Rush' in head
    by = {row[0]: dict(zip(head, row)) for row in rows[1:]}
    assert by['23']['In-Game Name'] == '\'=HYPERLINK("x")'
    assert by['21']['Furnace Level'] == 'FC10' and by['22']['Furnace Level'] == '30' and by['21']['Power (M)'] == '900.0'
    assert by['21']['Infantry Camp Level'] == 'FC5' and by['21']['Infantry Tier'] == 'T10'
    assert by['21']['Lancer Camp Level'] == 'FC3' and by['21']['Marksman Camp Level'] == ''
    assert by['21']['Marksman Tier'] == 'T11' and by['21']['Joiner Strength'] == str(5 + 10 + 3 + 9 + 0 + 11)
    assert by['23']['Joiner Strength'] == ''
    assert by['21']['Rally Leader'] == 'Yes' and by['22']['Rally Leader'] == 'No'
    for url in (f'/api/admin/tyrant/rounds/{rnd["id"]}/export', f'/api/admin/rounds/{rnd["id"]}/export'):
        r = client.get(url, headers=admin)
        assert r.status_code == 200 and r.mimetype.endswith('spreadsheetml.sheet')
        wb = openpyxl.load_workbook(io.BytesIO(r.data))
        ws = wb['Tyrant Poll Results']
        cells = {c.value: c for c in ws['B']}
        cell = cells['=HYPERLINK("x")']
        assert cell.data_type == 's' and cell.quotePrefix
        assert ws.freeze_panes == 'A2' and 'Summary' in wb.sheetnames
    assert client.get(f'/api/admin/tyrant/rounds/{rnd["id"]}/export.csv').status_code == 401


def test_delete_and_admin_edit(client, admin):
    rnd = _seed(client, admin)
    apps = client.get(f'/api/admin/tyrant/rounds/{rnd["id"]}/applications', headers=admin).json['applications']
    aid = next(a['id'] for a in apps if a['fid'] == '22')
    r = client.put(f'/api/admin/applications/{aid}', json={'answers': {'roles': ['event_prep']}}, headers=admin)
    assert r.status_code == 200 and r.json['answers']['roles'] == ['event_prep']
    r = client.put(f'/api/admin/applications/{aid}', json={'answers': {'roles': ['nope']}}, headers=admin)
    assert r.status_code == 400 and r.json['field'] == 'answers.roles'
    assert client.delete(f'/api/admin/applications/{aid}', headers=admin).json['deleted'] is True
    assert client.get(f'/api/admin/tyrant/rounds/{rnd["id"]}/summary', headers=admin).json['total'] == 2
    assert client.get('/api/profile/22').status_code == 200  # profile kept


def test_ministry_ignores_tyrant_profile_fields(client, admin):
    """Ministry keeps its free-form troops; the tyrant troop structure is only enforced on tyrant submits."""
    start_round(client, admin, 'M', event='ministry')
    r = client.put('/api/events/ministry/current/application/31',
                   json={'profile': {'game_name': 'M', 'alliance': 'MMM', 'troops': {'infantry': 5}}, 'answers': {}})
    assert r.status_code == 201 and r.json['profile']['troops'] == {'infantry': 5}


# ---------------------------------------------------------------- p2d: FC-only camps, filters, strength

def _full(camp, tier):
    return {k: {'furnace_level': camp, 'tier': tier} for k in ('infantry', 'lancer', 'marksman')}


def _seed_camps(client, admin):
    """41 = FC10 camps + T11 everywhere (the best joiner), 42 = FC10/T11 but FC9 lancer camp,
    43 = FC10 camps with T10 marksmen, 44 = LEGACY pre-FC camps (stored before the rule), 45 = no troop data."""
    rnd = start_round(client, admin, 'FDT camps', event=EVENT)
    assert put(client, '41', profile={'furnace_level': 'FC10', 'troops': _full('FC10', 11)}).status_code == 201
    t42 = _full('FC10', 11)
    t42['lancer'] = {'furnace_level': 'FC9', 'tier': 11}
    assert put(client, '42', profile={'furnace_level': 'FC10', 'troops': t42}).status_code == 201
    t43 = _full('FC10', 11)
    t43['marksman'] = {'furnace_level': 'FC10', 'tier': 10}
    assert put(client, '43', profile={'furnace_level': 'FC8', 'troops': t43}).status_code == 201
    # legacy row: the generic profile route keeps free-form troops and pre-FC codes (no data migration, owner rule)
    r = client.put('/api/admin/profiles/44', json={'game_name': 'Old', 'alliance': 'OLD', 'furnace_level': '25',
                                                  'troops': _full('25', 10)}, headers=admin)
    assert r.status_code in (200, 201), r.json
    assert put(client, '44', profile={'game_name': 'Old', 'alliance': 'OLD'}).status_code == 201
    assert put(client, '45').status_code == 201
    return rnd


def _fids(client, admin, rnd, qs):
    r = client.get(f'/api/admin/tyrant/rounds/{rnd["id"]}/applications?{qs}', headers=admin)
    assert r.status_code == 200, r.json
    return sorted(a['fid'] for a in r.json['applications'])


def test_tyrant_furnace_and_camps_fc_only_but_any_combination(client, admin):
    start_round(client, admin, 'FDT', event=EVENT)
    # any combination: camps above or below the furnace, T11 in an FC1 camp, T8 in an FC10 camp
    t = {'infantry': {'furnace_level': 'FC1', 'tier': 11}, 'lancer': {'furnace_level': 'fc10', 'tier': 8},
         'marksman': {'furnace_level': 'FC9', 'tier': 9}}
    r = put(client, '51', profile={'furnace_level': 'FC2', 'troops': t})
    assert r.status_code == 201, r.json
    assert r.json['profile']['troops']['lancer'] == {'furnace_level': 'FC10', 'tier': 8}  # normalised
    r = put(client, '51', profile={'furnace_level': '30'})
    assert r.status_code == 400 and r.json['code'] == 'VALIDATION_ERROR' and r.json['field'] == 'profile.furnace_level'
    r = put(client, '51', profile={'troops': {'infantry': {'furnace_level': '30', 'tier': 11}}})
    assert r.status_code == 400 and r.json['field'] == 'profile.troops.infantry.furnace_level'
    assert client.get('/api/profile/51').json['furnace_level'] == 'FC2'  # nothing written
    # admin edits go through the same rule
    aid = client.get('/api/admin/tyrant/rounds/current/applications', headers=admin).json['applications'][0]['id']
    r = client.put(f'/api/admin/applications/{aid}', json={'profile': {'furnace_level': '12'}}, headers=admin)
    assert r.status_code == 400 and r.json['field'] == 'profile.furnace_level'


def test_minister_keeps_pre_fc_levels(client, admin):
    start_round(client, admin, 'M', event='ministry')
    r = client.put('/api/events/ministry/current/application/61',
                   json={'profile': {'game_name': 'M', 'alliance': 'MMM', 'furnace_level': '30'}, 'answers': {}})
    assert r.status_code == 201 and r.json['profile']['furnace_level'] == '30'
    assert client.put('/api/profile/61', json={'furnace_level': '1'}).json['profile']['furnace_level'] == '1'


def test_legacy_pre_fc_values_shown_as_stored(client, admin):
    rnd = _seed_camps(client, admin)
    apps = {a['fid']: a for a in client.get(f'/api/admin/tyrant/rounds/{rnd["id"]}/applications',
                                             headers=admin).json['applications']}
    assert apps['44']['profile']['furnace_level'] == '25'
    assert apps['44']['profile']['troops']['infantry'] == {'furnace_level': '25', 'tier': 10}
    # resubmitting the legacy camp level is refused: the new rule applies on submit
    r = put(client, '44', profile={'troops': _full('25', 10)})
    assert r.status_code == 400 and r.json['field'] == 'profile.troops.infantry.furnace_level'


def test_admin_camp_and_tier_filters(client, admin):
    rnd = _seed_camps(client, admin)
    # 'FC10 camps with T11' in all three troop types: ONE query
    assert _fids(client, admin, rnd, 'min_camp=FC10&min_tier=T11') == ['41']
    assert _fids(client, admin, rnd, 'min_camp=fc10&min_tier=11&troop=all') == ['41']
    assert _fids(client, admin, rnd, 'min_tier=11') == ['41', '42']  # 'who has T11' everywhere
    assert _fids(client, admin, rnd, 'min_tier=11&troop=marksman') == ['41', '42']
    assert _fids(client, admin, rnd, 'min_tier=11&troop=infantry') == ['41', '42', '43']
    assert _fids(client, admin, rnd, 'min_camp=FC10&troop=lancer') == ['41', '43']
    assert _fids(client, admin, rnd, 'min_camp=FC9') == ['41', '42', '43']
    assert _fids(client, admin, rnd, 'min_camp=FC1') == ['41', '42', '43']  # legacy pre-FC camps never match
    assert _fids(client, admin, rnd, 'min_tier=10') == ['41', '42', '43', '44']
    assert _fids(client, admin, rnd, 'troop=lancer') == ['41', '42', '43', '44', '45']  # troop alone: no filter
    assert _fids(client, admin, rnd, 'min_tier=T11&min_furnace=FC10') == ['41', '42']  # combines with others
    base = f'/api/admin/tyrant/rounds/{rnd["id"]}/applications'
    for qs, field in (('min_camp=25', 'min_camp'), ('min_camp=FC11', 'min_camp'), ('min_tier=12', 'min_tier'),
                      ('min_tier=Tx', 'min_tier'), ('min_tier=0', 'min_tier'), ('troop=archer&min_tier=8', 'troop')):
        r = client.get(f'{base}?{qs}', headers=admin)
        assert r.status_code == 400 and r.json['code'] == 'VALIDATION_ERROR' and r.json['field'] == field, qs


def test_joiner_strength_sort(client, admin):
    rnd = _seed_camps(client, admin)
    r = client.get(f'/api/admin/tyrant/rounds/{rnd["id"]}/applications?sort=strength&dir=desc', headers=admin).json
    got = [(a['fid'], a['joiner_strength']) for a in r['applications']]
    # 41: 3*(10+11)=63; 42: 63-1=62; 43: 63-1=62; 44: pre-FC camps count 0 -> 3*10=30; 45: no data -> None, last
    assert got[0] == ('41', 63) and {f for f, _ in got[1:3]} == {'42', '43'} and got[1][1] == got[2][1] == 62
    assert got[3:] == [('44', 30), ('45', None)]
    r = client.get(f'/api/admin/tyrant/rounds/{rnd["id"]}/applications?sort=strength&dir=asc', headers=admin).json
    assert [a['fid'] for a in r['applications']][0] == '44' and r['applications'][-1]['fid'] == '45'  # blanks last


def test_summary_camp_level_counts(client, admin):
    rnd = _seed_camps(client, admin)
    s = client.get(f'/api/admin/tyrant/rounds/{rnd["id"]}/summary', headers=admin).json
    assert s['camp_levels']['infantry'] == {'FC10': 3, '25': 1, 'none': 1}
    assert s['camp_levels']['lancer'] == {'FC10': 2, 'FC9': 1, '25': 1, 'none': 1}
    assert list(s['camp_levels']['lancer']) == ['FC10', 'FC9', '25', 'none']  # FC high first, legacy, blanks last
    assert s['troop_tiers']['marksman'] == {'T11': 2, 'T10': 2, 'none': 1}


def test_exports_camp_columns(client, admin):
    rnd = _seed_camps(client, admin)
    r = client.get(f'/api/admin/tyrant/rounds/{rnd["id"]}/export.csv', headers=admin)
    rows = list(csv.reader(io.StringIO(r.data.decode('utf-8-sig'))))
    head = rows[0]
    for k in ('Infantry', 'Lancer', 'Marksman'):
        assert f'{k} Camp Level' in head and f'{k} Tier' in head
    assert head.index('Infantry Tier') == head.index('Infantry Camp Level') + 1
    by = {row[0]: dict(zip(head, row)) for row in rows[1:]}
    assert by['42']['Lancer Camp Level'] == 'FC9' and by['42']['Lancer Tier'] == 'T11'
    assert by['44']['Infantry Camp Level'] == '25'  # legacy shown as stored
    assert by['41']['Joiner Strength'] == '63' and by['45']['Joiner Strength'] == ''
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f'/api/admin/tyrant/rounds/{rnd["id"]}/export',
                                                      headers=admin).data))
    ws = wb['Tyrant Poll Results']
    xhead = [c.value for c in ws[1]]
    assert xhead == head
    col = xhead.index('Marksman Camp Level') + 1
    vals = {ws.cell(row=i, column=1).value: ws.cell(row=i, column=col).value for i in range(2, ws.max_row + 1)}
    assert vals['43'] == 'FC10' and vals['44'] == '25'
    summary = [[c.value for c in row] for row in wb['Summary'].iter_rows()]
    assert ['Lancer camp level', 'Players'] in summary and ['FC9', 1] in summary


def test_column_filters(client, admin):
    rnd = _seed(client, admin)  # 21 Zed AAA FC10 900M gems 50000 w1+w2 VC rally; 22 amy BBB '30' 100M no gems w3
    #                              no VC joiner+gathering; 23 AAA no stats gems 0 w1 VC no roles
    q = lambda qs: _fids(client, admin, rnd, qs)  # noqa: E731
    assert q('alliance=aaa,bbb') == ['21', '22', '23'] and q('alliance=bbb') == ['22']
    assert q('min_power=100000000&max_power=500000000') == ['22'] and q('min_power=1') == ['21', '22']
    assert q('min_gems=1') == ['21'] and q('max_gems=0') == ['23'] and q('min_gems=0') == ['21', '23']
    assert q('windows=w1') == ['21', '23'] and q('windows=w1,w2') == ['21'] and q('rush=1') == ['21', '23']
    assert q('vc=yes') == ['21', '23'] and q('vc=no') == ['22'] and q('vc=any') == ['21', '22', '23']
    assert q('roles=joiner,rally_leader') == ['21', '22'] and q('roles=joiner,gathering&roles_mode=all') == ['22']
    assert q('roles=joiner,rally_leader&roles_mode=all') == []
    assert q('min_furnace=FC1') == ['21']
    today = __import__('datetime').datetime.now(__import__('datetime').timezone.utc).strftime('%Y-%m-%d')
    assert q(f'submitted_from={today}&submitted_to={today}') == ['21', '22', '23']
    assert q('submitted_to=2000-01-01') == [] and q('days=1') == ['21', '22', '23']
    assert q('alliance=AAA&vc=yes&rush=1&min_gems=1') == ['21']  # AND
    # exact per-troop values (the summary chips), incl. 'none'
    assert q('infantry_camp=FC5') == ['21'] and q('infantry_camp=none') == ['23'] and q('marksman_tier=11') == ['21']
    assert q('infantry_tier=T8&infantry_camp=FC1') == ['22'] and q('lancer_tier=none') == ['22', '23']
    base = f'/api/admin/tyrant/rounds/{rnd["id"]}/applications'
    for qs, field in (('windows=w9', 'windows'), ('vc=maybe', 'vc'), ('rush=x', 'rush'), ('roles=boss', 'roles'),
                      ('roles=joiner&roles_mode=some', 'roles_mode'), ('submitted_from=2026-13-01', 'submitted_from'),
                      ('days=0', 'days'), ('min_power=-1', 'min_power'), ('max_gems=1.5', 'max_gems'),
                      ('infantry_camp=FC11', 'infantry_camp'), ('lancer_tier=12', 'lancer_tier')):
        r = client.get(f'{base}?{qs}', headers=admin)
        assert r.status_code == 400 and r.json['field'] == field, (qs, r.json)


def test_summary_and_exports_follow_filters(client, admin):
    rnd = _seed_camps(client, admin)
    s = client.get(f'/api/admin/tyrant/rounds/{rnd["id"]}/summary?lancer_camp=FC10', headers=admin).json
    assert s['total'] == 2 and s['round_total'] == 5 and s['filters'] == {'lancer_camp': 'FC10'}
    assert s['camp_levels']['lancer'] == {'FC10': 2} and s['troop_tiers']['marksman'] == {'T11': 1, 'T10': 1}
    s = client.get(f'/api/admin/tyrant/rounds/{rnd["id"]}/summary?min_camp=FC10&min_tier=11', headers=admin).json
    assert s['total'] == 1 and s['camp_levels']['infantry'] == {'FC10': 1}
    r = client.get(f'/api/admin/tyrant/rounds/{rnd["id"]}/export.csv?min_camp=FC10&min_tier=T11', headers=admin)
    rows = list(csv.reader(io.StringIO(r.data.decode('utf-8-sig'))))
    assert [row[0] for row in rows[1:]] == ['41']
    r = client.get(f'/api/admin/tyrant/rounds/{rnd["id"]}/export?min_tier=11', headers=admin)
    ws = openpyxl.load_workbook(io.BytesIO(r.data))['Tyrant Poll Results']
    assert sorted(ws.cell(row=i, column=1).value for i in range(2, ws.max_row + 1)) == ['41', '42']
    assert client.get(f'/api/admin/tyrant/rounds/{rnd["id"]}/export.csv?min_tier=99',
                      headers=admin).json['field'] == 'min_tier'
