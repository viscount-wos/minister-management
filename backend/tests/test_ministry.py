import io

import openpyxl
import pytest

from events.ministry import logic
from tests.conftest import apply, prefs, profile_id, start_round


def answers(c=0, r=0, t=0, g=0, fc=0, rfc=0, shards=0, slots=None):
    a = {'construction_speedups_days': c, 'research_speedups_days': r, 'troop_training_speedups_days': t,
         'general_speedups_days': g, 'fire_crystals': fc, 'refined_fire_crystals': rfc,
         'fire_crystal_shards': shards}
    if slots is not None:
        a['time_slots_by_day'] = slots
    return a


# ---------------------------------------------------------------- pure logic

def test_points_per_day():
    p = answers(c=2, r=3, t=40, g=1, fc=5, rfc=2, shards=7)
    assert logic.calculate_points(p, 'monday') == (2 + 1) * 1440 + 2 * 30000 + 5 * 2000
    assert logic.calculate_points(p, 'tuesday') == (3 + 1) * 1440 + 7 * 1000
    assert logic.calculate_points(p, 'friday') == (3 + 1) * 1440 + 7 * 1000  # research on friday
    assert logic.calculate_points(p, 'thursday') == 40
    assert logic.calculate_points(p, 'wednesday') == 0
    assert logic.calculate_points(answers(c=0.5), 'Monday') == 720


def test_slot_generation():
    ex = logic.generate_time_slots('exact_alignment')
    assert len(ex) == 48 and ex[0] == '00:00' and ex[1] == '00:30' and ex[-1] == '23:30'
    mx = logic.generate_time_slots('max_slots')
    assert len(mx) == 49 and mx[:3] == ['23:50', '00:20', '00:50'] and mx[-2:] == ['23:20', '23:50+']
    assert logic.matching_slots_for_pref(10, 'exact_alignment') == ['10:00', '10:30']
    assert logic.matching_slots_for_pref(0, 'max_slots') == ['23:50', '00:20', '00:50']
    assert logic.matching_slots_for_pref(23, 'max_slots') == ['22:50', '23:20', '23:50+']
    assert logic.slot_to_minutes('23:50') == -10 and logic.slot_to_minutes('23:50+') == 1430


def test_shared_slot_link():
    assert logic.get_shared_slot_link('monday', 'exact_alignment', 'tuesday') is None
    assert logic.get_shared_slot_link('monday', 'max_slots', 'tuesday') == ('tuesday', '23:50+', '23:50')
    assert logic.get_shared_slot_link('tuesday', 'max_slots', 'tuesday') == ('monday', '23:50', '23:50+')
    assert logic.get_shared_slot_link('monday', 'max_slots', 'friday') is None
    assert logic.get_shared_slot_link('thursday', 'max_slots', 'friday') == ('friday', '23:50+', '23:50')


# ---------------------------------------------------------------- auto-assign (HTTP)

def _assign(client, admin, day, ref='current'):
    r = client.post(f'/api/admin/ministry/rounds/{ref}/auto-assign', json={'day': day}, headers=admin)
    assert r.status_code == 200, r.json
    return r.json


def _slot_owner(result, slot):
    s = result['assignments'].get(slot) or []
    return s[0]['fid'] if s else None


def test_auto_assign_highest_points_wins_contested_slot(client, admin):
    start_round(client, admin, 'R')
    apply(client, '1', answers=answers(c=1, slots=prefs(['10:00'])), expect=201)
    apply(client, '2', answers=answers(c=5, slots=prefs(['10:00'])), expect=201)
    apply(client, '3', answers=answers(c=3, slots=prefs(['10:00'])), expect=201)
    apply(client, '4', answers=answers(c=9), expect=201)  # no time prefs
    res = _assign(client, admin, 'monday')
    assert _slot_owner(res, '10:00') == '2' and _slot_owner(res, '10:30') == '3'
    assert [u['fid'] for u in res['unassigned']] == ['4', '1']
    assert len(res['assignments']) == 48
    got = client.get('/api/admin/ministry/rounds/current/assignments/monday', headers=admin).json
    assert got['assignments']['10:00'][0]['fid'] == '2' and got['assignments']['10:00'][0]['points'] == 5 * 1440
    assert [u['fid'] for u in got['unassigned']] == ['4', '1']


def test_auto_assign_uses_day_type_prefs_and_research_friday(client, admin):
    start_round(client, admin, 'R', settings={'research_day': 'friday'})
    apply(client, '1', answers=answers(r=2, t=10, slots=prefs(['01:00'], ['05:00'], ['09:00'])), expect=201)
    assert _slot_owner(_assign(client, admin, 'friday'), '05:00') == '1'
    assert _slot_owner(_assign(client, admin, 'thursday'), '09:00') == '1'
    assert _slot_owner(_assign(client, admin, 'monday'), '01:00') == '1'
    r = client.post('/api/admin/ministry/rounds/current/auto-assign', json={'day': 'tuesday'}, headers=admin)
    assert r.status_code == 400 and r.json['code'] == 'VALIDATION_ERROR' and r.json['field'] == 'day'


def test_auto_assign_scoped_to_round(client, admin):
    r1 = start_round(client, admin, 'R1')
    apply(client, '1', answers=answers(c=1, slots=prefs(['10:00'])), expect=201)
    _assign(client, admin, 'monday')
    r2 = start_round(client, admin, 'R2')
    apply(client, '2', answers=answers(c=1, slots=prefs(['10:00'])), expect=201)
    res = _assign(client, admin, 'monday')
    assert _slot_owner(res, '10:00') == '2' and res['round_id'] == r2['id']
    old = client.get(f'/api/admin/ministry/rounds/{r1["id"]}/assignments/monday', headers=admin).json
    assert old['assignments']['10:00'][0]['fid'] == '1'


def test_max_slots_scheme_slots(client, admin):
    start_round(client, admin, 'R', settings={'time_slot_scheme': 'max_slots', 'research_day': 'friday'})
    apply(client, '1', answers=answers(c=5, slots=prefs(['10:00'])), expect=201)
    apply(client, '2', answers=answers(c=4, slots=prefs(['10:00'])), expect=201)
    apply(client, '3', answers=answers(c=3, slots=prefs(['10:00'])), expect=201)
    apply(client, '4', answers=answers(c=2, slots=prefs(['00:00'])), expect=201)
    res = _assign(client, admin, 'monday')
    assert len(res['assignments']) == 49
    assert [_slot_owner(res, s) for s in ('09:50', '10:20', '10:50')] == ['1', '2', '3']
    assert _slot_owner(res, '23:50') == '4'


def test_shared_2350_slot_combined_score_and_sync(client, admin):
    start_round(client, admin, 'R', settings={'time_slot_scheme': 'max_slots', 'research_day': 'tuesday'})
    # A: big monday, wants 23:00 monday. B: wants 00:00 tuesday, higher COMBINED score.
    apply(client, 'A'.replace('A', '11'), answers=answers(c=5, slots=prefs(['23:00'])), expect=201)
    apply(client, '22', answers=answers(c=3, r=4, slots=prefs([], ['00:00'])), expect=201)
    mon = _assign(client, admin, 'monday')
    assert _slot_owner(mon, '23:50+') == '22'
    tue = client.get('/api/admin/ministry/rounds/current/assignments/tuesday', headers=admin).json
    assert tue['assignments']['23:50'][0]['fid'] == '22'  # mirrored onto tuesday
    # Player 11 is still placed on monday elsewhere (22:50 / 23:20 via the 23:00 pref)
    assert _slot_owner(mon, '22:50') == '11'
    # Auto-assigning tuesday keeps the shared slot holder
    tue2 = _assign(client, admin, 'tuesday')
    assert _slot_owner(tue2, '23:50') == '22'
    mon2 = client.get('/api/admin/ministry/rounds/current/assignments/monday', headers=admin).json
    assert mon2['assignments']['23:50+'][0]['fid'] == '22'


def test_shared_slot_thursday_friday(client, admin):
    start_round(client, admin, 'R', settings={'time_slot_scheme': 'max_slots', 'research_day': 'friday'})
    apply(client, '1', answers=answers(t=10, slots=prefs([], [], ['23:00'])), expect=201)
    th = _assign(client, admin, 'thursday')
    assert _slot_owner(th, '23:50+') == '1'
    fr = client.get('/api/admin/ministry/rounds/current/assignments/friday', headers=admin).json
    assert fr['assignments']['23:50'][0]['fid'] == '1'


def test_no_shared_slot_in_exact_alignment(client, admin):
    start_round(client, admin, 'R')
    apply(client, '1', answers=answers(c=5, slots=prefs(['23:00'])), expect=201)
    _assign(client, admin, 'monday')
    tue = client.get('/api/admin/ministry/rounds/current/assignments/tuesday', headers=admin).json
    assert tue['assignments'] == {}


def test_manual_update_syncs_shared_slot(client, admin):
    start_round(client, admin, 'R', settings={'time_slot_scheme': 'max_slots'})
    apply(client, '1', answers=answers(c=1), expect=201)
    pid = profile_id(client, admin, '1')
    r = client.put('/api/admin/ministry/rounds/current/assignments/monday',
                   json={'assignments': {'23:50+': [{'player_id': pid}]}}, headers=admin)
    assert r.status_code == 200 and r.json['saved'] == 1
    tue = client.get('/api/admin/ministry/rounds/current/assignments/tuesday', headers=admin).json
    assert tue['assignments']['23:50'][0]['player_id'] == pid


def test_update_assignments_validation(client, admin):
    start_round(client, admin, 'R')
    apply(client, '1', expect=201)
    pid = profile_id(client, admin, '1')
    bad = [({'25:00': [{'player_id': pid}]}, 'assignments'),
           ({'10:00': [{'player_id': 999}]}, 'assignments'),
           ({'10:00': [{'player_id': pid}], '11:00': [{'player_id': pid}]}, 'assignments'),
           ({'23:50': [{'player_id': pid}]}, 'assignments')]  # max_slots slot under exact_alignment
    for body, field in bad:
        r = client.put('/api/admin/ministry/rounds/current/assignments/monday', json={'assignments': body},
                       headers=admin)
        assert r.status_code == 400 and r.json['field'] == field, body


def test_sticky_assignment_survives_auto_assign(client, admin):
    start_round(client, admin, 'R')
    apply(client, '1', answers=answers(c=1, slots=prefs(['10:00'])), expect=201)
    low = profile_id(client, admin, '1')
    apply(client, '2', answers=answers(c=9, slots=prefs(['05:00'])), expect=201)
    # admin pins the low scorer to 05:00 (the big player's preferred hour)
    client.put('/api/admin/ministry/rounds/current/assignments/monday',
               json={'assignments': {'05:00': [{'player_id': low, 'is_sticky': True}]}}, headers=admin)
    res = _assign(client, admin, 'monday')
    assert _slot_owner(res, '05:00') == '1' and res['assignments']['05:00'][0]['is_sticky'] is True
    assert _slot_owner(res, '05:30') == '2'
    got = client.get('/api/admin/ministry/rounds/current/assignments/monday', headers=admin).json
    assert got['assignments']['05:00'][0]['is_sticky'] is True
    res = _assign(client, admin, 'monday')  # still there on a second run
    assert _slot_owner(res, '05:00') == '1'


def test_scheme_change_remaps_assignments(client, admin):
    rnd = start_round(client, admin, 'R')
    apply(client, '1', answers=answers(c=5, slots=prefs(['10:00'])), expect=201)
    apply(client, '2', answers=answers(c=1, slots=prefs(['10:00'])), expect=201)
    _assign(client, admin, 'monday')  # 1 -> 10:00, 2 -> 10:30
    r = client.put(f'/api/admin/rounds/{rnd["id"]}', json={'settings': {'time_slot_scheme': 'max_slots'}},
                   headers=admin)
    assert r.status_code == 200 and r.json['remapped'] == 2
    got = client.get('/api/admin/ministry/rounds/current/assignments/monday', headers=admin).json
    # 10:00 -> 09:50 (nearest, 10 min), 10:30 -> 10:20 (10 min) per nearest-slot rule
    assert got['assignments']['09:50'][0]['fid'] == '1' and got['assignments']['10:20'][0]['fid'] == '2'


def test_remap_collision_higher_points_wins(client, admin):
    rnd = start_round(client, admin, 'R', settings={'time_slot_scheme': 'max_slots'})
    apply(client, '1', answers=answers(c=1), expect=201)
    a = profile_id(client, admin, '1')
    apply(client, '2', answers=answers(c=5), expect=201)
    b = profile_id(client, admin, '2')
    client.put('/api/admin/ministry/rounds/current/assignments/monday',
               json={'assignments': {'23:20': [{'player_id': a}], '23:50+': [{'player_id': b}]}}, headers=admin)
    # exact grid: both 23:20 and 23:50+ are nearest to 23:30 -> higher monday points (b) keeps it
    r = client.put(f'/api/admin/rounds/{rnd["id"]}', json={'settings': {'time_slot_scheme': 'exact_alignment'}},
                   headers=admin)
    assert r.json['remapped'] == 2  # monday b@23:30 + tuesday mirror b@00:00
    mon = client.get('/api/admin/ministry/rounds/current/assignments/monday', headers=admin).json
    assert {s: [p['fid'] for p in ps] for s, ps in mon['assignments'].items()} == {'23:30': ['2']}
    assert [u['fid'] for u in mon['unassigned']] == ['1']
    tue = client.get('/api/admin/ministry/rounds/current/assignments/tuesday', headers=admin).json
    assert tue['assignments']['00:00'][0]['fid'] == '2'


def test_publish_unpublish_and_public_schedule(client, admin):
    start_round(client, admin, 'R')
    apply(client, '1', name='Alice', alliance='AAA', answers=answers(c=1, slots=prefs(['10:00'])), expect=201)
    _assign(client, admin, 'monday')
    assert client.get('/api/events/ministry/current/schedule/monday').json['published'] is False
    r = client.post('/api/admin/ministry/rounds/current/publish', json={'day': 'thursday'}, headers=admin)
    r = client.post('/api/admin/ministry/rounds/current/publish', json={'day': 'monday'}, headers=admin)
    assert r.json['published_days'] == ['monday', 'thursday']  # weekday order
    assert client.get('/api/events/ministry/current/schedule').json['published_days'] == ['monday', 'thursday']
    s = client.get('/api/events/ministry/current/schedule/monday').json
    assert s['published'] and s['day_label'] == 'Monday - Construction'
    assert s['assignments'] == {'10:00': [{'game_name': 'Alice', 'alliance': 'AAA'}]}
    assert 'points' not in str(s)
    r = client.post('/api/admin/ministry/rounds/current/unpublish', json={'day': 'monday'}, headers=admin)
    assert r.json['published_days'] == ['thursday']
    r = client.post('/api/admin/ministry/rounds/current/publish', json={'day': 'friday'}, headers=admin)
    assert r.status_code == 400 and r.json['field'] == 'day'
    assert client.get('/api/events/ministry/current').json['settings']['published_days'] == ['thursday']


def test_player_own_assignments_and_heatmap(client, admin):
    start_round(client, admin, 'R')
    apply(client, '1', answers=answers(c=1, t=1, slots=prefs(['10:00'], [], ['10:00', '11:00'])), expect=201)
    apply(client, '2', answers=answers(slots=prefs(['10:00'])), expect=201)
    _assign(client, admin, 'monday')
    _assign(client, admin, 'thursday')
    r = client.get('/api/events/ministry/current/assignments/1').json
    assert r['assignments'] == {}  # drafts are not public (1c): only published days are listed
    for d in ('thursday', 'monday'):
        client.post('/api/admin/ministry/rounds/current/publish', json={'day': d}, headers=admin)
    r = client.get('/api/events/ministry/current/assignments/1').json
    assert r['assignments'] == {'monday': [{'time_slot': '10:00'}], 'thursday': [{'time_slot': '10:00'}]}
    assert client.get('/api/events/ministry/current/assignments/999').status_code == 404
    h = client.get('/api/events/ministry/current/heatmap').json
    assert h == {'construction': {'10:00': 2}, 'research': {}, 'troop': {'10:00': 1, '11:00': 1}}


def test_excel_export(client, admin):
    rnd = start_round(client, admin, 'Week of 13 Oct', settings={'research_day': 'friday'})
    apply(client, '1', name='Alice', alliance='AAA', answers=answers(c=1, slots=prefs(['10:00'])), expect=201)
    apply(client, '2', name='Bob', alliance='BBB', answers=answers(c=2), expect=201)
    _assign(client, admin, 'monday')
    for url in (f'/api/admin/rounds/{rnd["id"]}/export', '/api/admin/ministry/rounds/current/export'):
        r = client.get(url, headers=admin)
        assert r.status_code == 200
        assert r.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        assert 'minister_week_of_13_oct_' in r.headers['Content-Disposition']
        wb = openpyxl.load_workbook(io.BytesIO(r.data))
        assert wb.sheetnames == ['Monday - Construction', 'Friday - Research', 'Thursday - Troop Training',
                                 'Unassigned']
        mon = [[c.value for c in row] for row in wb['Monday - Construction'].iter_rows()]
        assert mon[0][0] == 'Time Slot' and mon[1][:4] == ['10:00', '1', 'AAA', 'Alice'] and mon[1][-1] == 1440
        assert ['UNASSIGNED PLAYERS'] == [r[0] for r in mon if r[0] == 'UNASSIGNED PLAYERS']
        assert mon[-1][:4] == ['Unassigned', '2', 'BBB', 'Bob']
        un = [[c.value for c in row] for row in wb['Unassigned'].iter_rows()]
        assert un[0] == ['Day', 'FID', 'Alliance', 'Game Name', 'Points']
        assert ['Monday - Construction', '2', 'BBB', 'Bob', 2880] in un


def test_export_unsupported_event(client, admin):
    rnd = start_round(client, admin, 'S', event='svs')
    r = client.get(f'/api/admin/rounds/{rnd["id"]}/export', headers=admin)
    assert r.status_code == 400 and r.json['code'] == 'EXPORT_NOT_SUPPORTED'


def test_export_json_import_roundtrip(client, admin):
    r1 = start_round(client, admin, 'R1')
    apply(client, '1', name='Alice', alliance='AAA', timezone='UTC',
          answers=answers(c=1, fc=3, slots=prefs(['10:00'], ['11:00'])), expect=201)
    data = client.get('/api/admin/ministry/rounds/current/export-json', headers=admin).json
    assert data['version'] == 2 and data['players'][0]['fire_crystals'] == 3
    r2 = start_round(client, admin, 'R2')
    data['players'].append({'fid': 'bad'})
    # v1.4-style entry with legacy list
    data['players'].append({'fid': '77', 'game_name': 'Legacy', 'alliance': 'old', 'time_slots': ['03:00'],
                            'avatar_image': 'http://x/a.png', 'stove_lv': 31})
    r = client.post(f'/api/admin/ministry/rounds/{r2["id"]}/import', json=data, headers=admin)
    assert r.status_code == 200
    assert (r.json['imported'], r.json['updated'], r.json['errors']) == (2, 0, 1)
    assert r.json['error_details'][0]['field'] == 'fid'
    apps = client.get(f'/api/admin/rounds/{r2["id"]}/applications', headers=admin).json['applications']
    by = {a['fid']: a for a in apps}
    assert by['1']['answers']['time_slots_by_day'] == prefs(['10:00'], ['11:00'])
    assert by['77']['answers']['time_slots_by_day'] == prefs(['03:00'], ['03:00'], ['03:00'])
    assert by['77']['profile']['stove_lv'] == 31 and by['77']['profile']['alliance'] == 'OLD'
    r = client.post(f'/api/admin/ministry/rounds/{r2["id"]}/import', json=data, headers=admin)
    assert (r.json['imported'], r.json['updated']) == (0, 2)
    assert client.post('/api/admin/ministry/rounds/current/import', json={'x': 1}, headers=admin).status_code == 400
    # r1 untouched
    assert len(client.get(f'/api/admin/rounds/{r1["id"]}/applications', headers=admin).json['applications']) == 1


def test_ministry_routes_reject_non_ministry_round(client, admin):
    rnd = start_round(client, admin, 'T', event='tyrant')
    r = client.get(f'/api/admin/ministry/rounds/{rnd["id"]}/assignments/monday', headers=admin)
    assert r.status_code == 404
    r = client.get('/api/admin/ministry/rounds/abc/assignments/monday', headers=admin)
    assert r.status_code == 404


@pytest.mark.parametrize('scheme', ['exact_alignment', 'max_slots'])
def test_every_player_placed_or_unassigned_exactly_once(client, admin, scheme):
    start_round(client, admin, 'R', settings={'time_slot_scheme': scheme})
    for i in range(1, 80):
        apply(client, str(i), answers=answers(c=i % 7, r=i % 5, slots=prefs([f'{i % 24:02d}:00'],
                                                                            [f'{(i * 3) % 24:02d}:00'])))
    for day in ('monday', 'tuesday', 'thursday'):
        _assign(client, admin, day)
        got = client.get(f'/api/admin/ministry/rounds/current/assignments/{day}', headers=admin).json
        placed = [p['fid'] for s in got['assignments'].values() for p in s]
        unplaced = [p['fid'] for p in got['unassigned']]
        assert len(placed) == len(set(placed))
        assert sorted(placed + unplaced, key=int) == [str(i) for i in range(1, 80)]
        assert all(len(s) == 1 for s in got['assignments'].values())
