"""Milestone 1c hardening: config refusal (H2/M4), login throttling (M3), CORS, Excel formula
injection (M5), export timestamps (L5), Unicode case folding (L7), busy DB -> 503 (L2),
per-entry import errors (L4), scheme-switch boundary re-sync (L3), public field minimisation (M6)."""
import io
import sqlite3

import openpyxl
import pytest

import app as app_module
from app import ConfigError, create_app
from core import auth
from tests.conftest import ADMIN_PW, MINISTER_PW, apply, prefs, profile_id, start_round

GOOD_KEY = 'a-long-random-test-secret-key-0123456789'


def _make(db_path, **cfg):
    base = {'DATABASE_PATH': db_path, 'SECRET_KEY': GOOD_KEY, 'ADMIN_PASSWORD': ADMIN_PW,
            'MINISTER_PASSWORD': MINISTER_PW, 'STATIC_DIR': '/nonexistent', 'DEV_MODE': False,
            'ALLOW_INSECURE_DEV': False}
    base.update(cfg)
    return create_app(base)


# ---------------------------------------------------------------- H2 / M4: config refusal

@pytest.mark.parametrize('key', ['', 'dev-secret-key', 'dev-secret-key-change-in-production', 'short-key'])
def test_production_refuses_missing_or_placeholder_secret(db_path, key):
    with pytest.raises(ConfigError):
        _make(db_path, SECRET_KEY=key)


def test_dev_placeholder_secret_needs_explicit_allow_insecure(db_path):
    with pytest.raises(ConfigError):
        _make(db_path, SECRET_KEY='dev-secret-key-change-in-production', DEV_MODE=True)
    app = _make(db_path, SECRET_KEY='dev-secret-key-change-in-production', DEV_MODE=True, ALLOW_INSECURE_DEV=True)
    assert app.config['SECRET_KEY'] == 'dev-secret-key-change-in-production'


def test_dev_missing_secret_gets_random_key(db_path):
    app = _make(db_path, SECRET_KEY='', DEV_MODE=True)
    assert len(app.config['SECRET_KEY']) >= 32


def test_placeholder_passwords_refused_outside_insecure_dev(db_path):
    with pytest.raises(ConfigError):
        _make(db_path, ADMIN_PASSWORD='admin123')
    with pytest.raises(ConfigError):
        _make(db_path, MINISTER_PASSWORD='minister123', DEV_MODE=True)
    _make(db_path, ADMIN_PASSWORD='admin123', DEV_MODE=True, ALLOW_INSECURE_DEV=True)


def test_env_defaults_only_with_allow_insecure_dev(monkeypatch):
    for k in ('ADMIN_PASSWORD', 'MINISTER_PASSWORD', 'ALLOW_INSECURE_DEV', 'MIGRATE_V14'):
        monkeypatch.delenv(k, raising=False)
    monkeypatch.setenv('FLASK_ENV', 'development')
    cfg = app_module._config_from_env()
    assert cfg['ADMIN_PASSWORD'] is None and cfg['MINISTER_PASSWORD'] is None and cfg['MIGRATE_V14'] is False
    monkeypatch.setenv('ALLOW_INSECURE_DEV', '1')
    monkeypatch.setenv('MIGRATE_V14', '1')
    cfg = app_module._config_from_env()
    assert cfg['ADMIN_PASSWORD'] == 'admin123' and cfg['ALLOW_INSECURE_DEV'] is True and cfg['MIGRATE_V14'] is True
    monkeypatch.setenv('FLASK_ENV', 'production')
    cfg = app_module._config_from_env()
    assert cfg['ADMIN_PASSWORD'] is None and cfg['ALLOW_INSECURE_DEV'] is False


def test_forged_placeholder_token_rejected_in_production(db_path):
    """M4: a token signed with the public compose placeholder key is worthless outside insecure dev."""
    from itsdangerous import URLSafeTimedSerializer
    forged = URLSafeTimedSerializer('dev-secret-key-change-in-production', salt=auth.TOKEN_SALT).dumps({'role': 'admin'})
    c = _make(db_path).test_client()
    assert c.get('/api/admin/me', headers={'Authorization': f'Bearer {forged}'}).status_code == 401


# ---------------------------------------------------------------- M3: login throttling

def test_limiter_unit_backoff_and_reset():
    now = [1000.0]
    lim = auth.LoginLimiter(free=3, base=2.0, cap=60, global_max=1000, clock=lambda: now[0])
    for _ in range(2):
        lim.failure('1.1.1.1')
        assert lim.retry_after('1.1.1.1') == 0
    lim.failure('1.1.1.1')                 # 3rd failure -> 2 s
    assert lim.retry_after('1.1.1.1') == pytest.approx(2.0)
    now[0] += 2.5
    assert lim.retry_after('1.1.1.1') == 0
    lim.failure('1.1.1.1')                 # 4th -> 4 s (exponential)
    assert lim.retry_after('1.1.1.1') == pytest.approx(4.0)
    assert lim.retry_after('2.2.2.2') == 0  # per IP
    for _ in range(10):
        now[0] += 100
        lim.failure('1.1.1.1')
    assert lim.retry_after('1.1.1.1') == pytest.approx(60)  # capped
    lim.success('1.1.1.1')
    assert lim.retry_after('1.1.1.1') == 0


def test_limiter_unit_global_budget():
    now = [0.0]
    lim = auth.LoginLimiter(free=100, global_window=60, global_max=5, clock=lambda: now[0])
    for i in range(5):
        lim.failure(f'10.0.0.{i}')
    assert lim.retry_after('10.9.9.9') == pytest.approx(60)  # a fresh IP waits too
    now[0] += 61
    assert lim.retry_after('10.9.9.9') == 0


def test_login_throttled_after_repeated_failures(client):
    for _ in range(auth.LOGIN_FREE_FAILURES):
        assert client.post('/api/admin/login', json={'password': 'wrong'}).status_code == 401
    r = client.post('/api/admin/login', json={'password': ADMIN_PW})  # even the right password waits
    assert r.status_code == 429 and r.json['code'] == 'TOO_MANY_ATTEMPTS'
    assert int(r.headers['Retry-After']) >= 1
    # another client address is unaffected (rightmost X-Forwarded-For entry = Cloud Run's)
    r = client.post('/api/admin/login', json={'password': ADMIN_PW},
                    headers={'X-Forwarded-For': '6.6.6.6, 203.0.113.9'})
    assert r.status_code == 200


def test_login_spoofed_forwarded_for_does_not_evade(client):
    for i in range(auth.LOGIN_FREE_FAILURES):
        client.post('/api/admin/login', json={'password': 'x'}, headers={'X-Forwarded-For': f'9.9.9.{i}, 198.51.100.7'})
    r = client.post('/api/admin/login', json={'password': 'x'}, headers={'X-Forwarded-For': '1.2.3.4, 198.51.100.7'})
    assert r.status_code == 429


def test_check_password_constant_time_shape(app):
    with app.app_context():
        assert auth._check_password(ADMIN_PW) == 'admin'
        assert auth._check_password(MINISTER_PW) == 'minister'
        assert auth._check_password(ADMIN_PW + 'x') is None
        app.config['MINISTER_PASSWORD'] = None
        assert auth._check_password('\x00disabled') is None  # a disabled role can never match


# ---------------------------------------------------------------- CORS

def test_cors_off_by_default_and_opt_in(db_path):
    c = _make(db_path).test_client()
    assert 'Access-Control-Allow-Origin' not in c.get('/api/events', headers={'Origin': 'https://evil.example'}).headers
    c = _make(db_path, CORS_ORIGINS=['https://ok.example']).test_client()
    assert c.get('/api/events', headers={'Origin': 'https://ok.example'}).headers.get(
        'Access-Control-Allow-Origin') == 'https://ok.example'
    assert 'Access-Control-Allow-Origin' not in c.get('/api/events', headers={'Origin': 'https://evil.example'}).headers


# ---------------------------------------------------------------- M5 / L5: exports

def test_excel_formula_injection_neutralised(client, admin):
    rnd = start_round(client, admin)
    names = {'1': '=HYPERLINK("http://x","y")', '2': '+SUM(1,2)', '3': '-2+3', '4': '@cmd', '5': 'Plain'}
    for fid, name in names.items():
        apply(client, fid, name=name, alliance='=A' if fid == '1' else 'ABC',
              answers={'construction_speedups_days': int(fid), 'time_slots_by_day': prefs(['10:00'])})
    client.post('/api/admin/ministry/rounds/current/auto-assign', json={'day': 'monday'}, headers=admin)
    r = client.get(f'/api/admin/rounds/{rnd["id"]}/export', headers=admin)
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    formulas = [(ws.title, c.coordinate) for ws in wb.worksheets for row in ws.iter_rows() for c in row
                if c.data_type == 'f']
    assert formulas == []
    seen = {c.value: c for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str)}
    for name in names.values():
        assert name in seen  # value preserved exactly
        if name[0] in '=+-@':
            assert seen[name].quotePrefix is True and seen[name].data_type == 's'
    assert '=A' in seen and seen['=A'].quotePrefix is True
    assert 'Unassigned' in wb.sheetnames


def test_export_json_timestamp_is_utc_iso(client, admin):
    start_round(client, admin)
    data = client.get('/api/admin/ministry/rounds/current/export-json', headers=admin).get_json(force=True)
    assert data['exported_at'].endswith('Z') and 'T' in data['exported_at']


def test_excel_unassigned_ties_in_player_id_order(client, admin):
    rnd = start_round(client, admin)
    for fid in ('30', '10', '20'):  # equal points, created in this order
        apply(client, fid, answers={'construction_speedups_days': 1})
    r = client.get(f'/api/admin/rounds/{rnd["id"]}/export', headers=admin)
    ws = openpyxl.load_workbook(io.BytesIO(r.data))['Monday - Construction']
    fids = [row[1] for row in ws.iter_rows(values_only=True) if row[0] == 'Unassigned']
    ids = [profile_id(client, admin, f) for f in fids]
    assert ids == sorted(ids)


# ---------------------------------------------------------------- L7: Unicode case folding

def test_unicode_casefold_filters(client, admin):
    rnd = start_round(client, admin)
    apply(client, '1', name='Élodie Straße', alliance='äbc')
    apply(client, '2', name='Bob', alliance='XYZ')
    r = client.get('/api/admin/profiles?alliance=ÄBC', headers=admin).json['profiles']
    assert [p['fid'] for p in r] == ['1']
    assert [p['fid'] for p in client.get('/api/admin/profiles?q=élodie', headers=admin).json['profiles']] == ['1']
    assert [p['fid'] for p in client.get('/api/admin/profiles?q=STRASSE', headers=admin).json['profiles']] == ['1']
    assert client.get('/api/admin/profiles?q=%25', headers=admin).json['profiles'] == []  # LIKE wildcard escaped
    apps = client.get(f'/api/admin/rounds/{rnd["id"]}/applications?alliance=äBC', headers=admin).json['applications']
    assert [a['fid'] for a in apps] == ['1']


# ---------------------------------------------------------------- L2: busy DB -> 503 RETRY

def test_locked_database_maps_to_503(db_path):
    app = _make(db_path, DB_BUSY_TIMEOUT=0.05)
    c = app.test_client()
    tok = c.post('/api/admin/login', json={'password': ADMIN_PW}).json['token']
    h = {'Authorization': f'Bearer {tok}'}
    rnd = c.post('/api/admin/events/ministry/start-new-round', json={'name': 'R'}, headers=h).json['round']
    other = sqlite3.connect(db_path)
    other.execute('BEGIN IMMEDIATE')
    try:
        r = c.post(f'/api/admin/ministry/rounds/{rnd["id"]}/publish', json={'day': 'monday'}, headers=h)
        assert r.status_code == 503 and r.json['code'] == 'RETRY' and r.headers['Retry-After']
        r = c.put(f'/api/admin/rounds/{rnd["id"]}', json={'name': 'x'}, headers=h)
        assert r.status_code == 503
    finally:
        other.rollback()
        other.close()
    assert c.post(f'/api/admin/ministry/rounds/{rnd["id"]}/publish', json={'day': 'monday'}, headers=h).status_code == 200


# ---------------------------------------------------------------- L4: import per-entry DB errors

def test_import_db_error_is_per_entry(client, admin, monkeypatch):
    start_round(client, admin)
    import core.applications as apps_mod
    real = apps_mod.save_application

    def flaky(round_row, profile_json, answers, commit=True):
        if profile_json['fid'] == '222':
            raise sqlite3.IntegrityError('UNIQUE constraint failed (simulated)')
        return real(round_row, profile_json, answers, commit=commit)
    monkeypatch.setattr(apps_mod, 'save_application', flaky)
    data = {'players': [{'fid': '111', 'game_name': 'A', 'alliance': 'AAA'},
                        {'fid': '222', 'game_name': 'B', 'alliance': 'BBB'},
                        {'fid': '333', 'game_name': 'C', 'alliance': 'CCC'}]}
    r = client.post('/api/admin/ministry/rounds/current/import', json=data, headers=admin)
    assert r.status_code == 200 and r.json['imported'] == 2 and r.json['errors'] == 1
    assert r.json['error_details'][0]['fid'] == '222'
    # the failed entry's profile insert was rolled back with its savepoint
    assert client.get('/api/profile/222').status_code == 404
    assert client.get('/api/profile/333').status_code == 200


# ---------------------------------------------------------------- L3: boundary re-sync after scheme switch

def test_scheme_switch_resyncs_shared_boundary(client, admin):
    rnd = start_round(client, admin, settings={'time_slot_scheme': 'exact_alignment', 'research_day': 'tuesday'})
    apply(client, '1', answers={'research_speedups_days': 5, 'time_slots_by_day': prefs(research=['00:00'])})
    apply(client, '2', answers={'construction_speedups_days': 5, 'time_slots_by_day': prefs(['10:00'])})
    pid1 = profile_id(client, admin, '1')
    pid2 = profile_id(client, admin, '2')
    client.put('/api/admin/ministry/rounds/current/assignments/tuesday',
               json={'assignments': {'00:00': [{'player_id': pid1}]}}, headers=admin)
    client.put('/api/admin/ministry/rounds/current/assignments/monday',
               json={'assignments': {'10:00': [{'player_id': pid2}]}}, headers=admin)
    r = client.put(f'/api/admin/rounds/{rnd["id"]}', json={'settings': {'time_slot_scheme': 'max_slots'}},
                   headers=admin)
    assert r.status_code == 200
    tue = client.get('/api/admin/ministry/rounds/current/assignments/tuesday', headers=admin).json
    mon = client.get('/api/admin/ministry/rounds/current/assignments/monday', headers=admin).json
    assert tue['assignments']['23:50'][0]['player_id'] == pid1
    assert mon['assignments']['23:50+'][0]['player_id'] == pid1  # mirrored: both days hold the same player
    assert [s for s, cs in mon['assignments'].items() if cs[0]['player_id'] == pid2] == ['09:50']


# ---------------------------------------------------------------- M6 / benign: response shapes

def test_public_profile_and_application_are_minimal(client, admin):
    start_round(client, admin)
    r = apply(client, '77', expect=201)
    assert set(r.json['profile']) == {'fid', 'game_name', 'alliance', 'timezone', 'furnace_level', 'power', 'troops',
                                      'discord_id', 'avatar_image', 'stove_lv', 'stove_lv_content'}
    assert set(r.json['application']) == {'fid', 'event', 'round_id', 'round_name', 'answers', 'updated_at'}
    assert set(client.get('/api/profile/77').json) == set(r.json['profile'])
    assert set(client.get('/api/events/ministry/current/application/77').json) == set(r.json['application'])
    r = client.put('/api/profile/77', json={'timezone': 'UTC'})
    assert 'id' not in r.json['profile'] and 'created_at' not in r.json['profile']
    # admin keeps the full shape
    full = client.get('/api/admin/profiles/77', headers=admin).json
    assert {'id', 'created_at', 'updated_at'} <= set(full)


def test_absent_strings_are_null_on_cards(client, admin):
    start_round(client, admin)
    apply(client, '5', answers={'time_slots_by_day': prefs(['10:00'])})
    res = client.post('/api/admin/ministry/rounds/current/auto-assign', json={'day': 'monday'}, headers=admin).json
    card = res['assignments']['10:00'][0]
    assert card['avatar_image'] is None and card['stove_lv_content'] is None and card['alliance'] == 'ABC'
    exp = client.get('/api/admin/ministry/rounds/current/export-json', headers=admin).get_json(force=True)
    assert exp['players'][0]['timezone'] is None and exp['players'][0]['avatar_image'] is None


def test_time_preferences_stored_sorted(client, admin):
    start_round(client, admin)
    apply(client, '6', answers={'time_slots_by_day': prefs(['15:00', '03:00', '10:00'])}, expect=201)
    got = client.get('/api/events/ministry/current/application/6').json['answers']['time_slots_by_day']
    assert got['construction'] == ['03:00', '10:00', '15:00']
