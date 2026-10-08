"""SVS event logic: EventSpec, admin list/summary and exports (modelled on events/tyrant/logic.py).

Profile (shared per FID with Frost Dragon Tyrant and Minister): game_name, alliance, troops (camp level + tier per
troop type). SVS does not ask the main furnace, power, gems or Discord ID. Round answers: hours, role, discord_vc,
language (events/svs/validation.py).
"""
import json
from collections import Counter
from datetime import datetime, timezone
from io import BytesIO

import openpyxl
from flask import Response
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.errors import validation_error
from core.exports import append_safe, to_csv_bytes
from core.troops import merge_troops, stored_troops
from events import EventSpec
from events.svs import filters as sf
from events.svs import validation as sv
from events.tyrant.logic import _furnace_sort, _slug, _tier_sort, _troop, joiner_strength, without_furnace
from events.tyrant.logic import round_applications as _tyrant_round_applications

EVENT = 'svs'
ROLE_LABELS = {'call': 'Call rallies', 'join': 'Join rallies'}
SORT_KEYS = ('submitted', 'updated', 'name', 'alliance', 'fid', 'strength', 'hours')


def round_settings(round_row):
    s = sv.default_settings()
    try:
        stored = json.loads(round_row['settings'] or '{}')
    except ValueError:
        stored = {}
    s.update({k: v for k, v in stored.items() if k in s})
    return s


def round_applications(db, round_row):
    """All applications of a round, newest first, with the CURRENT profile and joiner_strength (no furnace)."""
    return _tyrant_round_applications(db, round_row)


def filter_and_sort(apps, filters=None, sort='submitted', direction='desc'):
    apps = sf.apply_filters(apps, filters or {})
    if sort not in SORT_KEYS:
        raise validation_error('sort must be one of ' + ', '.join(SORT_KEYS), 'sort')
    if direction not in ('asc', 'desc'):
        raise validation_error('dir must be asc or desc', 'dir')
    keyfn = {
        'submitted': lambda x: (x['created_at'], x['id']),
        'updated': lambda x: (x['updated_at'], x['id']),
        'name': lambda x: (x['profile'].get('game_name') or '').casefold(),
        'alliance': lambda x: (x['profile'].get('alliance') or '').casefold(),
        'fid': lambda x: (len(x['fid'] or ''), x['fid'] or ''),
        'strength': lambda x: x.get('joiner_strength'),
        'hours': lambda x: len(x['answers'].get('hours') or []) or None,
    }[sort]
    present = [x for x in apps if keyfn(x) is not None]
    missing = [x for x in apps if keyfn(x) is None]  # blanks always last
    present.sort(key=keyfn, reverse=(direction == 'desc'))
    return present + missing


def summary(apps, settings):
    hours = sv.battle_hours(settings)
    hour_counts = Counter()
    roles = Counter()
    alliances = Counter()
    vc = 0
    troops = {k: Counter() for k in sv.TROOP_TYPES}
    camps = {k: Counter() for k in sv.TROOP_TYPES}
    for a in apps:
        ans, prof = a['answers'], a['profile']
        for h in ans.get('hours') or []:
            hour_counts[h] += 1
        roles[ans.get('role') or 'none'] += 1
        if ans.get('discord_vc'):
            vc += 1
        alliances[(prof.get('alliance') or '').strip().upper() or None] += 1
        for kind in sv.TROOP_TYPES:
            tier = _troop(prof, kind, 'tier')
            troops[kind][f'T{tier}' if tier else 'none'] += 1
            camps[kind][_troop(prof, kind, 'furnace_level') or 'none'] += 1
    return {
        'total': len(apps),
        'rally_callers': roles.get('call', 0),
        'joiners': roles.get('join', 0),
        'discord_vc': vc,
        'hours': [{'hour': h, 'count': hour_counts.get(h, 0)} for h in hours],
        'roles': {'call': roles.get('call', 0), 'join': roles.get('join', 0), 'none': roles.get('none', 0)},
        'alliances': [{'alliance': k, 'count': v}
                      for k, v in sorted(alliances.items(), key=lambda kv: (-kv[1], kv[0] or '~'))],
        'troop_tiers': {k: dict(sorted(c.items(), key=lambda kv: _tier_sort(kv[0]))) for k, c in troops.items()},
        'camp_levels': {k: dict(sorted(c.items(), key=lambda kv: _furnace_sort(kv[0]))) for k, c in camps.items()},
    }


# ---------------------------------------------------------------- exports

def _yes(v):
    return '' if v is None else ('Yes' if v else 'No')


def export_header(settings):
    return (['FID', 'In-Game Name', 'Alliance', 'Role']
            + [f'{h} UTC' for h in sv.battle_hours(settings)]
            + ['Discord VC']
            + [f'{k.capitalize()} {p}' for k in sv.TROOP_TYPES for p in ('Camp Level', 'Tier')]
            + ['Joiner Strength', 'Language', 'Submitted At (UTC)', 'Updated At (UTC)'])


def export_rows(apps, settings):
    hours = sv.battle_hours(settings)
    for a in apps:
        ans, prof = a['answers'], a['profile']
        mine = set(ans.get('hours') or [])
        row = [a['fid'], prof.get('game_name'), prof.get('alliance'), ROLE_LABELS.get(ans.get('role'), '')]
        row += ['Yes' if h in mine else 'No' for h in hours]
        row += [_yes(ans.get('discord_vc'))]
        for kind in sv.TROOP_TYPES:
            fl, tier = _troop(prof, kind, 'furnace_level'), _troop(prof, kind, 'tier')
            row += [fl, f'T{tier}' if tier else None]
        row += [a.get('joiner_strength'), ans.get('language'), a['created_at'], a['updated_at']]
        yield row


def build_csv(apps, settings):
    return to_csv_bytes(export_header(settings), export_rows(apps, settings))


def build_workbook(apps, settings, round_row):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'SVS Sign-ups'
    header = export_header(settings)
    ws.append(header)
    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='2B579A', end_color='2B579A', fill_type='solid')
    thin = Side(style='thin')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in ws[1]:
        c.font, c.fill, c.border = header_font, header_fill, border
        c.alignment = Alignment(horizontal='center')
    for row in export_rows(apps, settings):
        append_safe(ws, row)
        for c in ws[ws.max_row]:
            c.border = border
            c.alignment = Alignment(horizontal='center')
    for i in range(1, len(header) + 1):
        width = max(len(str(ws.cell(row=r, column=i).value or '')) for r in range(1, ws.max_row + 1))
        ws.column_dimensions[get_column_letter(i)].width = min(width + 4, 30)
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = 'A2'

    s = summary(apps, settings)
    sm = wb.create_sheet('Summary')
    append_safe(sm, ['Round', round_row['name']])
    sm.append(['Battle start (UTC)', settings['battle_start']])
    sm.append(['Battle hours', settings['battle_hours']])
    sm.append(['Total players', s['total']])
    sm.append(['Rally callers', s['rally_callers']])
    sm.append(['Joiners', s['joiners']])
    sm.append(['Discord VC', s['discord_vc']])
    sm.append([])
    sm.append(['Hour (UTC)', 'Players'])
    for h in s['hours']:
        sm.append([h['hour'], h['count']])
    sm.append([])
    sm.append(['Alliance', 'Players'])
    for a in s['alliances']:
        append_safe(sm, [a['alliance'] or '(none)', a['count']])
    for kind in sv.TROOP_TYPES:
        sm.append([])
        sm.append([f'{kind.capitalize()} camp level', 'Players'])
        for code, n in s['camp_levels'][kind].items():
            sm.append([code, n])
        sm.append([f'{kind.capitalize()} tier', 'Players'])
        for tier, n in s['troop_tiers'][kind].items():
            sm.append([tier, n])
    for col in ('A', 'B'):
        sm.column_dimensions[col].width = 28
    for c in sm['A']:
        c.font = Font(bold=True)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _download(data, mimetype, round_row, ext):
    filename = f'svs_{_slug(round_row["name"])}_{datetime.now(timezone.utc).strftime("%Y%m%d")}.{ext}'
    return Response(data, mimetype=mimetype, headers={'Content-Disposition': f'attachment; filename={filename}'})


def export_xlsx_response(db, round_row, filters=None):
    settings = round_settings(round_row)
    apps = sf.apply_filters(round_applications(db, round_row), filters or {})
    return _download(build_workbook(apps, settings, round_row),
                     'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', round_row, 'xlsx')


def export_csv_response(db, round_row, filters=None):
    settings = round_settings(round_row)
    apps = sf.apply_filters(round_applications(db, round_row), filters or {})
    return _download(build_csv(apps, settings), 'text/csv; charset=utf-8', round_row, 'csv')


# ---------------------------------------------------------------- EventSpec

class SvsEvent(EventSpec):
    key = EVENT
    has_rounds = True
    # Alliance: required for a NEW player only. upsert_profile checks the MERGED profile, so a known player's stored
    # alliance satisfies it and the wizard does not have to ask again (owner: "ask if it's a new player").
    required_profile_fields = ('game_name', 'alliance')
    # SVS asks none of these: an API client that sends them can't change the shared profile through SVS.
    ignored_profile_fields = ('furnace_level', 'power', 'discord_id', 'timezone')

    def default_settings(self):
        return sv.default_settings()

    def validate_settings(self, incoming, current):
        return sv.validate_settings(incoming, current)

    def public_settings(self, settings):
        return {'battle_start': settings.get('battle_start'), 'battle_hours': settings.get('battle_hours'),
                'hours': sv.battle_hours(settings)}

    def validate_profile(self, fields, existing=None, admin=False):
        if 'troops' in fields or not admin:
            sent = sv.validate_troops(fields.get('troops'), strict=not admin)
            fields['troops'] = merge_troops(stored_troops(existing), sent)
        return fields

    def validate_answers(self, answers, round_, existing=None, admin=False):
        return sv.validate_answers(answers, round_settings(round_), existing=existing, strict=not admin)

    def decorate_application(self, app, round_):
        app['joiner_strength'] = joiner_strength(app.get('profile') or {})
        return without_furnace(app)

    def export_round(self, round_):
        from core.db import get_db
        return export_xlsx_response(get_db(), round_)
