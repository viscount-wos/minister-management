"""Ministry logic ported from v1.4 (points, slot schemes, shared 23:50 slot, auto-assign,
remapping, export), now scoped to a round.

Players here are dicts flattened from (application answers + profile):
``id``/``player_id`` = profiles.id (what ministry_assignments.player_id stores).
"""
import json
import logging
import re
import sqlite3
from datetime import datetime, timezone
from io import BytesIO

import openpyxl
from flask import Response
from openpyxl.styles import Alignment, Font, PatternFill

from events import EventSpec
from events.ministry import validation as mv

logger = logging.getLogger(__name__)

WEEKDAY_ORDER = ['monday', 'tuesday', 'wednesday', 'thursday', 'friday', 'saturday', 'sunday']
DAY_LABELS = {
    'monday': 'Monday - Construction',
    'tuesday': 'Tuesday - Research',
    'friday': 'Friday - Research',
    'thursday': 'Thursday - Troop Training',
}


# ---------------------------------------------------------------- points

def calculate_points(player, day):
    """v1.4 maths. 1 day = 1440 min.

    monday:  (construction + general) * 1440 + refined*30000 + fire_crystals*2000
    tuesday/friday (research): (research + general) * 1440 + shards*1000
    thursday: troop_training days (1 point per day)
    """
    def n(key):
        return player.get(key) or 0

    construction_mins = n('construction_speedups_days') * 24 * 60
    research_mins = n('research_speedups_days') * 24 * 60
    troop_days = n('troop_training_speedups_days')
    general_mins = n('general_speedups_days') * 24 * 60
    d = day.lower()
    if d == 'monday':
        return int(construction_mins + general_mins + n('refined_fire_crystals') * 30000 + n('fire_crystals') * 2000)
    if d in ('tuesday', 'friday'):
        return int(research_mins + general_mins + n('fire_crystal_shards') * 1000)
    if d == 'thursday':
        return int(troop_days)
    return 0


# ---------------------------------------------------------------- days & slots

MINISTRY_WEEKDAYS = ('monday', 'tuesday', 'thursday', 'friday')  # every day that can ever be a ministry day


def valid_days(research_day):
    return ['monday', research_day, 'thursday']


def active_published_days(settings):
    """Published days that are active days of the round (week order). Stale entries, e.g. an old
    research day left over in stored settings, are never served publicly."""
    valid = valid_days(settings.get('research_day', 'tuesday'))
    return sort_days_by_week([d for d in settings.get('published_days') or [] if d in valid])


def sort_days_by_week(days):
    return sorted(days, key=lambda d: (WEEKDAY_ORDER.index(d) if d in WEEKDAY_ORDER else len(WEEKDAY_ORDER), d))


def day_type_for(day, research_day):
    if day == 'monday':
        return 'construction'
    if day == 'thursday':
        return 'troop'
    if day == research_day:
        return 'research'
    return 'construction'


def generate_time_slots(scheme):
    """'exact_alignment' -> 00:00 .. 23:30 (48). 'max_slots' -> 23:50, 00:20 .. 23:20, 23:50+ (49)."""
    if scheme == 'exact_alignment':
        return [f'{h:02d}:{m:02d}' for h in range(24) for m in (0, 30)]
    slots = ['23:50']
    hour, minute = 0, 20
    while True:
        slot = f'{hour:02d}:{minute:02d}'
        if slot == '23:50':
            slots.append('23:50+')
            break
        slots.append(slot)
        minute += 30
        if minute >= 60:
            minute -= 60
            hour += 1
    return slots


def matching_slots_for_pref(hour, scheme):
    if scheme == 'exact_alignment':
        return [f'{hour:02d}:00', f'{hour:02d}:30']
    prev_h = (hour - 1) % 24
    result = ['23:50'] if hour == 0 else [f'{prev_h:02d}:50']
    result.append(f'{hour:02d}:20')
    result.append('23:50+' if hour == 23 else f'{hour:02d}:50')
    return result


def slot_to_minutes(slot):
    if slot == '23:50':
        return -10
    if slot == '23:50+':
        return 1430
    return int(slot[:2]) * 60 + int(slot[3:5])


def get_shared_slot_link(day, scheme, research_day):
    """(other_day, this_slot, other_slot) for the shared 23:50 boundary, or None.

    max_slots only: monday 23:50+ == tuesday 23:50 (research tuesday);
    thursday 23:50+ == friday 23:50 (research friday).
    """
    if scheme != 'max_slots':
        return None
    pairs = []
    if research_day == 'tuesday':
        pairs.append(('monday', 'tuesday'))
    if research_day == 'friday':
        pairs.append(('thursday', 'friday'))
    for earlier, later in pairs:
        if day == earlier:
            return (later, '23:50+', '23:50')
        if day == later:
            return (earlier, '23:50', '23:50+')
    return None


def compute_shared_winner(players, earlier_day, later_day, research_day):
    earlier_type = day_type_for(earlier_day, research_day)
    later_type = day_type_for(later_day, research_day)

    def is_candidate(p):
        by_day = p.get('time_slots_by_day', {})
        return ('23:00' in by_day.get(earlier_type, [])) or ('00:00' in by_day.get(later_type, []))

    candidates = [p for p in players if is_candidate(p)]
    if not candidates:
        return None
    for p in candidates:
        p['_combined'] = calculate_points(p, earlier_day) + calculate_points(p, later_day)
    candidates.sort(key=lambda p: p['_combined'], reverse=True)
    return candidates[0]


# ---------------------------------------------------------------- data access

def round_settings(round_row):
    s = dict(mv.DEFAULT_SETTINGS)
    s.update(json.loads(round_row['settings'] or '{}'))
    return s


def get_round_players(db, round_id):
    """All applicants of a round as flat player dicts (newest application first, like v1.4)."""
    rows = db.execute(
        'SELECT a.id AS application_id, a.answers, a.created_at AS applied_at, a.updated_at AS app_updated_at, '
        'p.* FROM applications a JOIN profiles p ON p.id = a.player_id '
        'WHERE a.round_id = ? ORDER BY a.created_at DESC, a.id ASC', (round_id,)).fetchall()
    return [_flatten(r) for r in rows]


def get_round_player(db, round_id, player_id):
    r = db.execute(
        'SELECT a.id AS application_id, a.answers, a.created_at AS applied_at, a.updated_at AS app_updated_at, '
        'p.* FROM applications a JOIN profiles p ON p.id = a.player_id '
        'WHERE a.round_id = ? AND a.player_id = ?', (round_id, player_id)).fetchone()
    return _flatten(r) if r else None


def _flatten(r):
    d = dict(r)
    answers = json.loads(d.pop('answers') or '{}')
    p = {
        'id': d['id'],
        'player_id': d['id'],
        'application_id': d['application_id'],
        'fid': d['fid'],
        'game_name': d['game_name'],
        # absent strings are null (v1.4 mixed '' and null; docs/API.md convention is null)
        'alliance': d.get('alliance') or None,
        'timezone': d.get('timezone') or None,
        'furnace_level': d.get('furnace_level') or None,
        'avatar_image': d.get('avatar_image') or None,
        'stove_lv': d.get('stove_lv'),
        'stove_lv_content': d.get('stove_lv_content') or None,
        'created_at': d['applied_at'],
        'updated_at': d['app_updated_at'],
    }
    for f in mv.NUMERIC_FIELDS:
        p[f] = answers.get(f) or 0
    by_day = answers.get('time_slots_by_day') or {}
    p['time_slots_by_day'] = {t: list(by_day.get(t, [])) for t in mv.DAY_TYPES}
    p['time_slots'] = sorted({s for t in mv.DAY_TYPES for s in p['time_slots_by_day'][t]})
    return p


def card(player, day, prefs=None, sticky=False):
    """Assignment card shape (same keys as v1.4)."""
    return {
        'id': player['id'],
        'player_id': player['id'],
        'fid': player['fid'],
        'game_name': player['game_name'],
        'points': calculate_points(player, day),
        'preferred_times': sorted(prefs) if prefs is not None else [],
        'avatar_image': player.get('avatar_image') or None,
        'stove_lv': player.get('stove_lv'),
        'stove_lv_content': player.get('stove_lv_content') or None,
        'alliance': player.get('alliance') or None,
        'furnace_level': player.get('furnace_level') or None,
        'is_sticky': bool(sticky),
    }


def _insert_assignment(cursor, round_id, player_id, day, slot, position, is_assigned, is_sticky):
    cursor.execute(
        'INSERT INTO ministry_assignments (round_id, player_id, day, time_slot, position, is_assigned, is_sticky, '
        "created_at) VALUES (?, ?, ?, ?, ?, ?, ?, strftime('%Y-%m-%dT%H:%M:%SZ','now'))",
        (round_id, player_id, day, slot, position, 1 if is_assigned else 0, 1 if is_sticky else 0))


def sync_shared_boundary(cursor, round_id, day, scheme, research_day):
    """Mirror ``day``'s 23:50 boundary slot onto the linked day (same real time). Caller commits."""
    link = get_shared_slot_link(day, scheme, research_day)
    if not link:
        return
    other_day, this_slot, other_slot = link
    row = cursor.execute('SELECT player_id, is_sticky FROM ministry_assignments '
                         'WHERE round_id = ? AND day = ? AND time_slot = ?', (round_id, day, this_slot)).fetchone()
    cursor.execute('DELETE FROM ministry_assignments WHERE round_id = ? AND day = ? AND time_slot = ?',
                   (round_id, other_day, other_slot))
    if row:
        pid = row['player_id']
        cursor.execute('DELETE FROM ministry_assignments WHERE round_id = ? AND day = ? AND player_id = ?',
                       (round_id, other_day, pid))
        _insert_assignment(cursor, round_id, pid, other_day, other_slot, 0, True, row['is_sticky'])


# ---------------------------------------------------------------- auto-assign

def auto_assign(db, round_row, day):
    settings = round_settings(round_row)
    research_day = settings['research_day']
    scheme = settings['time_slot_scheme']
    round_id = round_row['id']
    day = mv.validate_day(day, research_day)
    day_type = day_type_for(day, research_day)

    players = get_round_players(db, round_id)
    for p in players:
        p['points'] = calculate_points(p, day)
    players.sort(key=lambda p: p['points'], reverse=True)  # stable: ties keep newest-first order
    by_id = {p['id']: p for p in players}

    time_slots = generate_time_slots(scheme)
    cursor = db.cursor()

    # Sticky placements survive (only for players still in the round)
    sticky_slots = {}
    sticky_ids = set()
    for row in cursor.execute('SELECT player_id, time_slot FROM ministry_assignments '
                              'WHERE round_id = ? AND day = ? AND is_sticky = 1', (round_id, day)).fetchall():
        p = by_id.get(row['player_id'])
        if p is None or row['time_slot'] not in time_slots:
            continue  # gone from the round, or a slot from another scheme: re-place normally
        sticky_slots[row['time_slot']] = card(p, day, p['time_slots_by_day'].get(day_type, []), sticky=True)
        sticky_ids.add(p['id'])

    assignments = {slot: [] for slot in time_slots}
    unassigned = []
    for slot, c in sticky_slots.items():
        if slot in assignments:
            assignments[slot].append(c)

    shared_locked = set()
    link = get_shared_slot_link(day, scheme, research_day)
    if link:
        other_day, this_slot, _ = link
        earlier, later = (day, other_day) if this_slot == '23:50+' else (other_day, day)
        if this_slot in assignments and not assignments[this_slot]:
            winner = compute_shared_winner(players, earlier, later, research_day)
            if winner and winner['id'] not in sticky_ids:
                assignments[this_slot] = [card(winner, day, winner['time_slots_by_day'].get(day_type, []))]
                shared_locked.add(winner['id'])

    for p in players:
        if p['id'] in sticky_ids or p['id'] in shared_locked:
            continue
        prefs = p['time_slots_by_day'].get(day_type, [])
        matching = []
        # sorted: v1.4 iterated a set here, which made slot choice depend on hash order
        for pref in sorted(set(prefs)):
            if ':' in pref:
                matching.extend(matching_slots_for_pref(int(pref.split(':')[0]), scheme))
        placed = False
        for slot in matching:
            if slot in assignments and not assignments[slot]:
                assignments[slot].append(card(p, day, prefs))
                placed = True
                break
        if not placed:
            unassigned.append(card(p, day, prefs))

    cursor.execute('DELETE FROM ministry_assignments WHERE round_id = ? AND day = ?', (round_id, day))
    for slot, slot_players in assignments.items():
        for pos, c in enumerate(slot_players):
            _insert_assignment(cursor, round_id, c['id'], day, slot, pos, True, c['is_sticky'])
    sync_shared_boundary(cursor, round_id, day, scheme, research_day)
    db.commit()
    logger.info('Auto-assign round=%s day=%s assigned=%s unassigned=%s', round_id, day,
                sum(1 for s in assignments.values() if s), len(unassigned))
    return {'day': day, 'assignments': assignments, 'unassigned': unassigned}


def get_assignments(db, round_row, day):
    settings = round_settings(round_row)
    research_day = settings['research_day']
    day = mv.validate_day(day, research_day)
    day_type = day_type_for(day, research_day)
    players = {p['id']: p for p in get_round_players(db, round_row['id'])}
    rows = db.execute('SELECT * FROM ministry_assignments WHERE round_id = ? AND day = ? '
                      'ORDER BY time_slot, position', (round_row['id'], day)).fetchall()
    assignments = {}
    assigned_ids = set()
    for r in rows:
        p = players.get(r['player_id'])
        if p is None:
            continue  # application deleted
        c = card(p, day, p['time_slots_by_day'].get(day_type, []), sticky=r['is_sticky'])
        c['assignment_id'] = r['id']
        c['position'] = r['position']
        c['is_assigned'] = bool(r['is_assigned'])
        assignments.setdefault(r['time_slot'], []).append(c)
        assigned_ids.add(p['id'])
    unassigned = [card(p, day, p['time_slots_by_day'].get(day_type, []))
                  for pid, p in players.items() if pid not in assigned_ids]
    unassigned.sort(key=lambda x: x['points'], reverse=True)
    return {'day': day, 'assignments': assignments, 'unassigned': unassigned}


def update_assignments(db, round_row, day, assignments):
    """Replace a day's assignments (drag-and-drop save). Max one player per slot, as v1.4."""
    from core.errors import validation_error
    settings = round_settings(round_row)
    research_day = settings['research_day']
    scheme = settings['time_slot_scheme']
    day = mv.validate_day(day, research_day)
    if not isinstance(assignments, dict):
        raise validation_error('assignments must be an object of slot -> [player]', 'assignments')
    valid_slots = set(generate_time_slots(scheme))
    member_ids = {r['player_id'] for r in db.execute('SELECT player_id FROM applications WHERE round_id = ?',
                                                     (round_row['id'],))}
    rows = []
    seen = set()
    for slot, slot_players in assignments.items():
        if not slot_players:
            continue
        if slot not in valid_slots:
            raise validation_error(f'Invalid time slot {slot!r} for scheme {scheme}', 'assignments')
        if not isinstance(slot_players, list) or not isinstance(slot_players[0], dict):
            raise validation_error(f'assignments[{slot}] must be a list of player objects', 'assignments')
        p = slot_players[0]
        pid = p.get('player_id', p.get('id'))
        if not isinstance(pid, int) or pid not in member_ids:
            raise validation_error(f'Player {pid!r} has no application in this round', 'assignments')
        if pid in seen:
            raise validation_error(f'Player {pid} is assigned to more than one slot', 'assignments')
        seen.add(pid)
        rows.append((slot, pid, p.get('is_assigned', True), bool(p.get('is_sticky'))))
    cursor = db.cursor()
    cursor.execute('DELETE FROM ministry_assignments WHERE round_id = ? AND day = ?', (round_row['id'], day))
    for slot, pid, is_assigned, sticky in rows:
        _insert_assignment(cursor, round_row['id'], pid, day, slot, 0, is_assigned, sticky)
    sync_shared_boundary(cursor, round_row['id'], day, scheme, research_day)
    db.commit()
    return {'day': day, 'saved': len(rows)}


def remap_assignments_between_schemes(db, round_id, new_scheme):
    """Move every assignment of the round to the nearest slot of the new scheme.

    Collisions: higher points wins, loser becomes unassigned. Returns kept count. Caller commits.
    """
    target = [(s, slot_to_minutes(s)) for s in generate_time_slots(new_scheme)]
    rows = [dict(r) for r in db.execute('SELECT player_id, day, time_slot, is_sticky FROM ministry_assignments '
                                        'WHERE round_id = ?', (round_id,)).fetchall()]
    if not rows:
        return 0
    players = {p['id']: p for p in get_round_players(db, round_id)}
    by_day = {}
    for r in rows:
        by_day.setdefault(r['day'], []).append(r)
    kept = 0
    for day, day_rows in by_day.items():
        planned = []
        for r in day_rows:
            src = slot_to_minutes(r['time_slot'])
            nearest = min(target, key=lambda t: abs(t[1] - src))[0]
            p = players.get(r['player_id'])
            planned.append({'row': r, 'target': nearest, 'points': calculate_points(p, day) if p else 0})
        planned.sort(key=lambda x: x['points'], reverse=True)
        taken = set()
        db.execute('DELETE FROM ministry_assignments WHERE round_id = ? AND day = ?', (round_id, day))
        for pl in planned:
            if pl['target'] in taken:
                continue
            taken.add(pl['target'])
            _insert_assignment(db, round_id, pl['row']['player_id'], day, pl['target'], 0, True,
                               pl['row']['is_sticky'])
            kept += 1
    resync_shared_boundaries(db, round_id, new_scheme)
    return kept


def resync_shared_boundaries(db, round_id, scheme):
    """After a scheme switch (L3), make both sides of each shared 23:50 boundary hold the same player.

    The remap moves each day independently, so e.g. exact->max can fill Tuesday 23:50 (from 00:00)
    while Monday 23:50+ stays empty. Mirror from the earlier day if its boundary slot is occupied,
    otherwise from the later day. Caller commits.
    """
    if scheme != 'max_slots':
        return
    rnd = db.execute('SELECT settings FROM rounds WHERE id = ?', (round_id,)).fetchone()
    research_day = json.loads(rnd['settings'] or '{}').get('research_day', 'tuesday') if rnd else 'tuesday'
    for earlier in ('monday', 'thursday'):
        link = get_shared_slot_link(earlier, scheme, research_day)
        if not link or link[1] != '23:50+':
            continue
        later = link[0]

        def occupied(day, slot):
            return db.execute('SELECT 1 FROM ministry_assignments WHERE round_id = ? AND day = ? AND time_slot = ?',
                              (round_id, day, slot)).fetchone() is not None
        if occupied(earlier, '23:50+'):
            sync_shared_boundary(db, round_id, earlier, scheme, research_day)
        elif occupied(later, '23:50'):
            sync_shared_boundary(db, round_id, later, scheme, research_day)


# ---------------------------------------------------------------- heatmap / public

def heatmap(db, round_id):
    result = {t: {} for t in mv.DAY_TYPES}
    for p in get_round_players(db, round_id):
        for t in mv.DAY_TYPES:
            for slot in p['time_slots_by_day'][t]:
                result[t][slot] = result[t].get(slot, 0) + 1
    return result


def published_schedule(db, round_row, day):
    settings = round_settings(round_row)
    day = str(day).lower()
    if day not in active_published_days(settings):
        return {'published': False, 'day': day}
    rows = db.execute(
        'SELECT a.time_slot, p.game_name, p.alliance FROM ministry_assignments a '
        'JOIN profiles p ON p.id = a.player_id '
        'JOIN applications ap ON ap.player_id = a.player_id AND ap.round_id = a.round_id '
        'WHERE a.round_id = ? AND a.day = ? AND a.is_assigned = 1 ORDER BY a.time_slot, a.position',
        (round_row['id'], day)).fetchall()
    assignments = {}
    for r in rows:
        assignments.setdefault(r['time_slot'], []).append({'game_name': r['game_name'],
                                                           'alliance': r['alliance'] or ''})
    return {'published': True, 'day': day, 'day_label': DAY_LABELS.get(day, day), 'assignments': assignments}


def player_assignments(db, round_id, player_id):
    rows = db.execute('SELECT day, time_slot FROM ministry_assignments WHERE round_id = ? AND player_id = ? '
                      'AND is_assigned = 1 ORDER BY time_slot', (round_id, player_id)).fetchall()
    out = {}
    for r in rows:
        out.setdefault(r['day'], []).append({'time_slot': r['time_slot']})
    return {d: out[d] for d in sort_days_by_week(out)}


# ---------------------------------------------------------------- export

EXPORT_HEADERS = ['Time Slot', 'FID', 'Alliance', 'Game Name', 'Construction (days)', 'Research (days)',
                  'Troop Training (days)', 'General (days)', 'Fire Crystals', 'Refined Fire Crystals',
                  'Crystal Shards', 'Points']
EXPORT_WIDTHS = [15, 15, 10, 25, 18, 15, 20, 15, 13, 18, 14, 12]


FORMULA_TRIGGERS = ('=', '+', '-', '@', '\t', '\r')


def _append_safe(ws, row):
    """Append a row; neutralise user text that a spreadsheet could run as a formula (M5).

    Strings starting with = + - @ TAB or CR are forced to plain text cells with Excel's quote
    prefix, so the value is shown exactly as typed and never evaluated (also after an edit).
    """
    ws.append(row)
    for cell in ws[ws.max_row]:
        v = cell.value
        if isinstance(v, str) and v[:1] in FORMULA_TRIGGERS:
            cell.data_type = 's'
            cell.quotePrefix = True


def _player_row(slot_label, p, points):
    return [slot_label, p['fid'], p.get('alliance', ''), p['game_name'],
            p['construction_speedups_days'], p['research_speedups_days'], p['troop_training_speedups_days'],
            p['general_speedups_days'], p['fire_crystals'], p['refined_fire_crystals'], p['fire_crystal_shards'],
            points]


def build_workbook(db, round_row):
    """v1.4 layout: one sheet per day (assigned rows, then an UNASSIGNED section), plus an
    'Unassigned' sheet summarising who is unplaced on each day."""
    settings = round_settings(round_row)
    research_day = settings['research_day']
    players = get_round_players(db, round_row['id'])
    by_id = {p['id']: p for p in players}
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')
    sep_fill = PatternFill(start_color='F4B942', end_color='F4B942', fill_type='solid')
    sep_font = Font(bold=True, color='000000')
    research_label = 'Tuesday - Research' if research_day == 'tuesday' else 'Friday - Research'
    days = [('monday', 'Monday - Construction'), (research_day, research_label),
            ('thursday', 'Thursday - Troop Training')]
    unassigned_summary = []

    def style_header(ws):
        for c in ws[1]:
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal='center')

    for day_key, title in days:
        ws = wb.create_sheet(title=title)
        ws.append(EXPORT_HEADERS)
        style_header(ws)
        assigned_ids = set()
        for r in db.execute('SELECT player_id, time_slot FROM ministry_assignments WHERE round_id = ? AND day = ? '
                            'AND is_assigned = 1 ORDER BY time_slot, position',
                            (round_row['id'], day_key)).fetchall():
            p = by_id.get(r['player_id'])
            if p is None:
                continue
            assigned_ids.add(p['id'])
            label = '23:50 (+1d)' if r['time_slot'] == '23:50+' else r['time_slot']
            _append_safe(ws, _player_row(label, p, calculate_points(p, day_key)))
        unassigned = [p for p in players if p['id'] not in assigned_ids]
        # points DESC, ties by player id ASC (v1.4's Excel order)
        unassigned.sort(key=lambda p: (-calculate_points(p, day_key), p['id']))
        if unassigned:
            sep_row = ws.max_row + 2
            ws.append([])
            ws.append(['UNASSIGNED PLAYERS'] + [''] * (len(EXPORT_HEADERS) - 1))
            for c in ws[sep_row]:
                c.fill = sep_fill
                c.font = sep_font
            for p in unassigned:
                pts = calculate_points(p, day_key)
                _append_safe(ws, _player_row('Unassigned', p, pts))
                unassigned_summary.append([title, p['fid'], p.get('alliance', ''), p['game_name'], pts])
        for i, w in enumerate(EXPORT_WIDTHS):
            ws.column_dimensions[chr(65 + i)].width = w

    ws = wb.create_sheet(title='Unassigned')
    ws.append(['Day', 'FID', 'Alliance', 'Game Name', 'Points'])
    style_header(ws)
    for row in unassigned_summary:
        _append_safe(ws, row)
    for i, w in enumerate([28, 15, 10, 25, 12]):
        ws.column_dimensions[chr(65 + i)].width = w
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _slug(text):
    return re.sub(r'[^A-Za-z0-9]+', '_', text).strip('_').lower()[:40] or 'round'


# ---------------------------------------------------------------- JSON export / import

def export_json(db, round_row):
    players = []
    for p in get_round_players(db, round_row['id']):
        entry = {k: p[k] for k in ('fid', 'game_name', 'alliance', 'avatar_image', 'stove_lv',
                                   'stove_lv_content', 'timezone', 'time_slots_by_day')}
        for f in mv.NUMERIC_FIELDS:
            entry[f] = p[f]
        players.append(entry)
    from core.validation import now_iso
    return {'version': 2, 'exported_at': now_iso(),
            'round': {'id': round_row['id'], 'name': round_row['name'], 'settings': round_settings(round_row)},
            'players': players}


def import_json(db, round_row, data):
    """Upsert profiles + applications in the round from a v1 (v1.4) or v2 export. Admin only."""
    from core.applications import get_application, save_application
    from core.errors import validation_error
    from core.profiles import resolve_fid_for_write, upsert_profile, validate_profile_fields
    from core.errors import ApiError

    if not isinstance(data, dict) or not isinstance(data.get('players'), list):
        raise validation_error('Invalid format: expected {players: [...]}', 'players')
    imported = updated = errors = 0
    error_list = []
    if not db.in_transaction:
        db.execute('BEGIN')  # one transaction; each entry in its own savepoint (L4)
    for i, p in enumerate(data['players']):
        db.execute('SAVEPOINT import_entry')
        try:
            if not isinstance(p, dict):
                raise validation_error('player entry must be an object')
            fid, existing = resolve_fid_for_write(p.get('fid'))
            fields = validate_profile_fields({'game_name': p.get('game_name') or 'Unknown',
                                              'alliance': p.get('alliance') or None,
                                              'timezone': p.get('timezone') or None}, field_prefix='',
                                             existing=existing)
            answers = {f: p.get(f, 0) for f in mv.NUMERIC_FIELDS}
            if 'time_slots_by_day' in p:
                answers['time_slots_by_day'] = p['time_slots_by_day']
            elif 'time_slots' in p:
                answers['time_slots'] = p['time_slots']
            prev = get_application(round_row['id'], existing['id']) if existing else None
            answers = mv.validate_answers(answers, existing=json.loads(prev['answers'] or '{}') if prev else None)
            existed = existing is not None
            profile, _ = upsert_profile(fid, fields, commit=False, existing=existing)
            # keep legacy avatar/stove data on brand-new profiles imported from a v1.4 backup
            if not existed and any(p.get(k) for k in ('avatar_image', 'stove_lv', 'stove_lv_content')):
                db.execute('UPDATE profiles SET avatar_image=?, stove_lv=?, stove_lv_content=? WHERE id=?',
                           (p.get('avatar_image') or None, p.get('stove_lv'), p.get('stove_lv_content') or None,
                            profile['id']))
            _, created = save_application(round_row, profile, answers, commit=False)
            if created:
                imported += 1
            else:
                updated += 1
            db.execute('RELEASE import_entry')
        except (ApiError, sqlite3.Error) as e:
            db.execute('ROLLBACK TO import_entry')
            db.execute('RELEASE import_entry')
            errors += 1
            if isinstance(e, ApiError):
                msg, field = e.message, e.field
            else:
                logger.warning('import entry %s failed: %s', i, e)
                msg, field = 'Database rejected this entry (e.g. duplicate FID)', None
            error_list.append({'index': i, 'fid': p.get('fid') if isinstance(p, dict) else None,
                               'error': msg, 'field': field})
    db.commit()
    return {'imported': imported, 'updated': updated, 'errors': errors, 'error_details': error_list[:50]}


# ---------------------------------------------------------------- EventSpec

class MinistryEvent(EventSpec):
    key = 'ministry'
    has_rounds = True
    required_profile_fields = ('game_name', 'alliance')  # v1.4 required both

    def default_settings(self):
        return json.loads(json.dumps(mv.DEFAULT_SETTINGS))

    def validate_settings(self, incoming, current):
        return mv.validate_settings(incoming, current, valid_days)

    def carry_over_settings(self, previous):
        s = dict(previous)
        s['published_days'] = []  # a new round starts unpublished
        s.pop('legacy_closing_time_raw', None)
        return s

    def public_settings(self, settings):
        s = {k: settings.get(k) for k in mv.DEFAULT_SETTINGS}
        s['published_days'] = active_published_days(s)
        return s

    def validate_answers(self, answers, round_, existing=None):
        return mv.validate_answers(answers, existing=existing)

    def decorate_application(self, app, round_):
        rd = round_settings(round_)['research_day']
        a = app['answers']
        app['monday_points'] = calculate_points(a, 'monday')
        app['research_points'] = calculate_points(a, rd)
        app['thursday_points'] = calculate_points(a, 'thursday')
        app['research_day'] = rd
        return app

    def on_settings_changed(self, db, round_, old, new):
        if old.get('time_slot_scheme') != new.get('time_slot_scheme'):
            return {'remapped': remap_assignments_between_schemes(db, round_['id'], new['time_slot_scheme'])}
        return {}

    def on_application_deleted(self, db, round_id, player_id):
        db.execute('DELETE FROM ministry_assignments WHERE round_id = ? AND player_id = ?', (round_id, player_id))

    def export_round(self, round_):
        from core.db import get_db
        data = build_workbook(get_db(), round_)
        filename = f'minister_{_slug(round_["name"])}_{datetime.now(timezone.utc).strftime("%Y%m%d")}.xlsx'
        return Response(data, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                        headers={'Content-Disposition': f'attachment; filename={filename}'})
