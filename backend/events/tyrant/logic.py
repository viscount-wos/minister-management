"""Frost Dragon Tyrant event logic: EventSpec, admin list/summary and exports.

Profile (shared, per FID): game_name, alliance, discord_id, power, troops (per troop type: CAMP level + tier).
Tyrant does NOT ask the main furnace (owner decision p2e): a tyrant submit ignores ``profile.furnace_level``, and
tyrant admin rows, summary and exports carry no furnace (the shared profile keeps it for Minister).
Round answers: availability (window ids of the round), discord_vc, gem_spend, roles, language.
"""
import json
from collections import Counter
from datetime import datetime, timezone
from io import BytesIO

import openpyxl
from flask import Response
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.exports import append_safe, to_csv_bytes
from core.furnace import FURNACE_LEVELS, fc_number, furnace_ordinal
from core.troops import merge_troops, stored_troops
from core.validation import validate_number
from core.errors import validation_error
from events import EventSpec
from events.tyrant import filters as tf
from events.tyrant import validation as tv

EVENT = 'tyrant'
ROLE_LABELS = {'rally_leader': 'Rally Leader', 'joiner': 'Joiner', 'gathering': 'Gathering/Looting',
               'battle_mgmt': 'Battle Management', 'event_prep': 'Event Preparation'}
SORT_KEYS = ('submitted', 'updated', 'name', 'alliance', 'fid', 'power', 'gems', 'strength')


def round_settings(round_row):
    s = tv.default_settings()
    try:
        stored = json.loads(round_row['settings'] or '{}')
    except ValueError:
        stored = {}
    if isinstance(stored.get('windows'), list):
        s['windows'] = stored['windows']
    return s


def _slug(text):
    import re
    return re.sub(r'[^A-Za-z0-9]+', '_', text).strip('_').lower()[:40] or 'round'


def round_applications(db, round_row):
    """All applications of a round, newest first, each with the CURRENT profile (as ministry)."""
    from core.applications import application_to_json
    from core.profiles import profile_to_json
    rows = db.execute('SELECT a.*, p.fid, p.id AS pid, p.game_name AS p_game_name, p.alliance AS p_alliance, '
                      'p.timezone AS p_timezone, p.power AS p_power, '
                      'p.troops AS p_troops, p.discord_id AS p_discord_id, p.avatar_image AS p_avatar_image, '
                      'p.stove_lv AS p_stove_lv, p.stove_lv_content AS p_stove_lv_content, '
                      'p.created_at AS p_created_at, p.updated_at AS p_updated_at '
                      'FROM applications a JOIN profiles p ON p.id = a.player_id WHERE a.round_id = ? '
                      'ORDER BY a.created_at DESC, a.id DESC', (round_row['id'],)).fetchall()
    out = []
    for r in rows:
        app = application_to_json(r, round_row)
        prof = {k[2:]: r[k] for k in r.keys() if k.startswith('p_')}
        prof['id'] = r['pid']
        prof['fid'] = r['fid']
        app['profile'] = profile_to_json(prof)
        app['joiner_strength'] = joiner_strength(app['profile'])
        out.append(without_furnace(app))
    return out


def without_furnace(app):
    """Tyrant rows carry no main furnace (owner decision p2e: camp levels + tiers are the strength signal). The
    shared profile keeps ``furnace_level`` for Minister; it is only left out of tyrant admin rows/exports/MCP."""
    for key in ('profile', 'profile_snapshot'):
        if isinstance(app.get(key), dict):
            app[key].pop('furnace_level', None)
    return app


def _troops(profile):
    t = profile.get('troops')
    return t if isinstance(t, dict) else {}


def _troop(profile, kind, key):
    entry = _troops(profile).get(kind)
    if isinstance(entry, dict):
        v = entry.get(key)
        if key == 'furnace_level':
            return v if isinstance(v, str) and v in FURNACE_LEVELS else None
        return v if isinstance(v, int) and not isinstance(v, bool) else None
    return None


def joiner_strength(profile):
    """Joiner strength (admin sort key, owner rule p2d): the sum over infantry, lancer and marksman of
    camp FC number (FC1=1 .. FC10=10; pre-FC or blank = 0) + tier number (T8=8 .. T11=11; blank = 0).
    0..63; 63 = FC10 camps with T11 troops in all three. None when no camp level or tier is filled in at all
    (sorted last). A sum, not a min: it ranks a player with one weak camp below an otherwise equal one, but
    still above a player who is weak everywhere."""
    total, any_value = 0, False
    for kind in tv.TROOP_TYPES:
        camp, tier = _troop(profile, kind, 'furnace_level'), _troop(profile, kind, 'tier')
        if camp or tier:
            any_value = True
        total += fc_number(camp) + (tier or 0)
    return total if any_value else None


def filter_and_sort(apps, filters=None, settings=None, sort='submitted', direction='desc'):
    """``filters``: the dict of events.tyrant.filters.parse_filters (all AND); then sort (blanks always last)."""
    apps = tf.apply_filters(apps, filters or {}, settings or {})
    if sort not in SORT_KEYS:
        raise validation_error('sort must be one of ' + ', '.join(SORT_KEYS), 'sort')
    if direction not in ('asc', 'desc'):
        raise validation_error('dir must be asc or desc', 'dir')
    keyfn = {
        'submitted': lambda x: (x['created_at'], x['id']),
        'updated': lambda x: (x['updated_at'], x['id']),
        'name': lambda x: (x['profile'].get('game_name') or '').casefold(),
        'alliance': lambda x: (x['profile'].get('alliance') or '').casefold(),
        'fid': lambda x: (len(x['fid'] or ''), x['fid'] or ''),
        'power': lambda x: x['profile'].get('power'),
        'gems': lambda x: x['answers'].get('gem_spend'),
        'strength': lambda x: x.get('joiner_strength'),
    }[sort]
    present = [x for x in apps if keyfn(x) is not None]
    missing = [x for x in apps if keyfn(x) is None]  # blanks always last
    present.sort(key=keyfn, reverse=(direction == 'desc'))
    return present + missing


def summary(apps, settings):
    windows = settings.get('windows') or []
    win_counts = Counter()
    rush_ids = {w['id'] for w in windows if w.get('rush')}
    opening_rush = vc = 0
    roles = Counter()
    alliances = Counter()
    troops = {k: Counter() for k in tv.TROOP_TYPES}
    camps = {k: Counter() for k in tv.TROOP_TYPES}
    for a in apps:
        ans, prof = a['answers'], a['profile']
        avail = set(ans.get('availability') or [])
        for w in avail:
            win_counts[w] += 1
        if avail & rush_ids:
            opening_rush += 1
        if ans.get('discord_vc'):
            vc += 1
        for r in ans.get('roles') or []:
            roles[r] += 1
        alliances[(prof.get('alliance') or '').strip().upper() or None] += 1
        for kind in tv.TROOP_TYPES:
            tier = _troop(prof, kind, 'tier')
            troops[kind][f'T{tier}' if tier else 'none'] += 1
            camps[kind][_troop(prof, kind, 'furnace_level') or 'none'] += 1
    return {
        'total': len(apps),
        'opening_rush': opening_rush,
        'discord_vc': vc,
        'windows': [dict(w, count=win_counts.get(w['id'], 0)) for w in windows],
        'alliances': [{'alliance': k, 'count': v}
                      for k, v in sorted(alliances.items(), key=lambda kv: (-kv[1], kv[0] or '~'))],
        'roles': {r: roles.get(r, 0) for r in tv.ROLES},
        'troop_tiers': {k: dict(sorted(c.items(), key=lambda kv: _tier_sort(kv[0]))) for k, c in troops.items()},
        # camp level per troop type: FC10 first, legacy pre-FC codes (shown as stored) after FC1, 'none' last
        'camp_levels': {k: dict(sorted(c.items(), key=lambda kv: _furnace_sort(kv[0]))) for k, c in camps.items()},
    }


def _tier_sort(label):
    return 999 if label == 'none' else -int(label[1:])


def _furnace_sort(label):
    return 999 if label == 'none' else -furnace_ordinal(label)


# ---------------------------------------------------------------- exports

def _window_label(w):
    return f'{w["start"]}-{w["end"]}' + (' Opening Rush' if w.get('rush') else '')


def export_header(settings):
    return (['FID', 'In-Game Name', 'Alliance', 'Discord ID']
            + [_window_label(w) for w in settings['windows']]
            + ['Discord VC', 'Power (M)', 'Est. Max Gem Spend']
            + [f'{k.capitalize()} {p}' for k in tv.TROOP_TYPES for p in ('Camp Level', 'Tier')]
            + ['Joiner Strength']
            + [ROLE_LABELS[r] for r in tv.ROLES]
            + ['Language', 'Submitted At (UTC)', 'Updated At (UTC)'])


def export_rows(apps, settings):
    yes = lambda b: 'Yes' if b else 'No'  # noqa: E731 (tyrantpoll wrote Yes/No)
    for a in apps:
        ans, prof = a['answers'], a['profile']
        avail = set(ans.get('availability') or [])
        power = prof.get('power')
        row = [a['fid'], prof.get('game_name'), prof.get('alliance'), prof.get('discord_id')]
        row += [yes(w['id'] in avail) for w in settings['windows']]
        row += [yes(ans.get('discord_vc')),
                round(power / 1_000_000, 2) if power is not None else None, ans.get('gem_spend')]
        for kind in tv.TROOP_TYPES:
            fl, tier = _troop(prof, kind, 'furnace_level'), _troop(prof, kind, 'tier')
            row += [fl, f'T{tier}' if tier else None]
        row += [a.get('joiner_strength')]
        roles = set(ans.get('roles') or [])
        row += [yes(r in roles) for r in tv.ROLES]
        row += [ans.get('language'), a['created_at'], a['updated_at']]
        yield row


def build_csv(apps, settings):
    return to_csv_bytes(export_header(settings), export_rows(apps, settings))


def build_workbook(apps, settings, round_row):
    """tyrantpoll's styled sheet (blue header, borders, filter, frozen header) + a Summary sheet."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Tyrant Poll Results'
    header = export_header(settings)
    ws.append(header)
    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='2B579A', end_color='2B579A', fill_type='solid')
    thin = Side(style='thin')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in ws[1]:
        c.font, c.fill, c.border = header_font, header_fill, border
        c.alignment = Alignment(horizontal='center')
    for row in export_rows(apps, settings):
        append_safe(ws, row)
        for c in ws[ws.max_row]:
            c.border = border
            c.alignment = Alignment(horizontal='center')
    for i in range(1, len(header) + 1):
        width = max(len(str(ws.cell(row=r, column=i).value or '')) for r in range(1, ws.max_row + 1))
        ws.column_dimensions[get_column_letter(i)].width = min(width + 4, 30)
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = 'A2'

    s = summary(apps, settings)
    sm = wb.create_sheet('Summary')
    append_safe(sm, ['Round', round_row['name']])
    sm.append(['Total players', s['total']])
    sm.append(['Opening rush', s['opening_rush']])
    sm.append(['Discord VC', s['discord_vc']])
    sm.append([])
    sm.append(['Window (UTC)', 'Players'])
    for w in s['windows']:
        sm.append([_window_label(w), w['count']])
    sm.append([])
    sm.append(['Role', 'Players'])
    for r, n in s['roles'].items():
        sm.append([ROLE_LABELS[r], n])
    sm.append([])
    sm.append(['Alliance', 'Players'])
    for a in s['alliances']:
        append_safe(sm, [a['alliance'] or '(none)', a['count']])
    for kind in tv.TROOP_TYPES:
        sm.append([])
        sm.append([f'{kind.capitalize()} camp level', 'Players'])
        for code, n in s['camp_levels'][kind].items():
            sm.append([code, n])
        sm.append([f'{kind.capitalize()} tier', 'Players'])
        for tier, n in s['troop_tiers'][kind].items():
            sm.append([tier, n])
    for col in ('A', 'B'):
        sm.column_dimensions[col].width = 28
    for c in sm['A']:
        c.font = Font(bold=True)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _download(data, mimetype, round_row, ext):
    filename = f'tyrant_{_slug(round_row["name"])}_{datetime.now(timezone.utc).strftime("%Y%m%d")}.{ext}'
    return Response(data, mimetype=mimetype, headers={'Content-Disposition': f'attachment; filename={filename}'})


def export_xlsx_response(db, round_row, filters=None):
    settings = round_settings(round_row)
    apps = tf.apply_filters(round_applications(db, round_row), filters or {}, settings)
    return _download(build_workbook(apps, round_settings(round_row), round_row),
                     'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', round_row, 'xlsx')


def export_csv_response(db, round_row, filters=None):
    settings = round_settings(round_row)
    apps = tf.apply_filters(round_applications(db, round_row), filters or {}, settings)
    return _download(build_csv(apps, round_settings(round_row)), 'text/csv; charset=utf-8', round_row, 'csv')


# ---------------------------------------------------------------- EventSpec

class TyrantEvent(EventSpec):
    key = EVENT
    has_rounds = True
    required_profile_fields = ('game_name', 'alliance')  # tyrantpoll required name, FID and alliance
    ignored_profile_fields = ('furnace_level',)  # owner decision p2e: Tyrant does not ask the main furnace

    def default_settings(self):
        return tv.default_settings()

    def validate_settings(self, incoming, current):
        return tv.validate_settings(incoming, current)

    def public_settings(self, settings):
        return {'windows': settings.get('windows') or []}

    def validate_profile(self, fields, existing=None, admin=False):
        if 'troops' in fields:
            # shared merge rule (core/troops.py): a blank never clears what SVS (or an earlier sign-up) stored
            fields['troops'] = merge_troops(stored_troops(existing), tv.validate_troops(fields['troops']))
        if fields.get('power') is not None:
            # power arrives as an absolute number (the UI converts "millions"); cap at 10^13 is core's
            fields['power'] = validate_number(fields['power'], 'profile.power', maximum=10 ** 13, integer=True)
        return fields

    def validate_answers(self, answers, round_, existing=None, admin=False):
        return tv.validate_answers(answers, round_settings(round_), existing=existing)

    def decorate_application(self, app, round_):
        return without_furnace(app)

    def export_round(self, round_):
        from core.db import get_db
        return export_xlsx_response(get_db(), round_)
